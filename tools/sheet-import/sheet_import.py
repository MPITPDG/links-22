#!/usr/bin/env python3
"""
sheet_import.py - one-time importer: the Bajaj monthly Google workbook (.xlsx download) -> Export > Sheet view
tables on LinksDB20 (TBL_EXPT_SHEET_MONTH / _ROW / _GRIDROW / _TABMETA / _LOG).

It never connects to a database. It reads the workbook and writes, into --out-dir (outside any git work tree):
    IMPORT_EXPORT_SHEET_<YYYYMM>_LINKSDB20_<date>.sql            the import (review it, then run it in SSMS)
    ROLLBACK_IMPORT_EXPORT_SHEET_<YYYYMM>_LINKSDB20_<date>.sql   removes exactly that import (refuses after edits)
    CHECK_FFJOB_EXPORT_SHEET_<YYYYMM>.sql                        read-only: which FF JOB values match a Links job
    IMPORT_EXPORT_SHEET_<YYYYMM>_REPORT.txt                      what was read, mapped, skipped and why
The generated files hold customer shipment data: never commit them.

Usage (Python 3.8+, openpyxl 3.x; run isolated, the workbook is untrusted data):
    python -I -X utf8 tools/sheet-import/sheet_import.py <workbook.xlsx> --month 2026-10 --title "Bajaj October Sheet 2026"
           --out-dir <folder outside the repo> [--cmpcode 30] [--citycode MUM] [--skip-tab "<tab>"] [--date YYYYMMDD]
    python -I -X utf8 tools/sheet-import/sheet_import.py --selftest

Sources of truth it follows (see README.md):
    src/app/export/sheet/sheet-registry.ts           SHEET_COLUMNS, LAYOUT_A / LAYOUT_B, DEFAULT_TABS, OPTIONAL_WO_TABS, GRID_TABS
    _DB_FIX/EXPORT_SHEET_VIEW_LINKSDB20_20261021.sql table DDL, the 10 Colour CODE seeds, the canonical CELL_FMT form
                                                     (USP_EXPT_SHEET_SAVE_FMT), the date rule (USP_EXPT_SHEET_SAVE_APPLY)
    Links_v5_API ExportSheetController               XmlSafe (WO values), Json.NET Formatting.None + GuardJson (DATA / TABS / META)

The "Column Access" tab is never read (it holds staff e-mail addresses); "Colour CODE" is only compared with the seeded
legend; "KPI Summary" is computed by the Sheet view and is skipped.
"""
import argparse
import colorsys
import datetime
import decimal
import hashlib
import os
import re
import sys
import warnings
import xml.etree.ElementTree as ET

VERSION = '1.0 (2026-10-21)'
IMPORT_USER = 'SHEET IMPORT'          # MAKER_ID / MODIFY_ID / LOG USER_ID + USER_NAME of everything this import writes
SORT_STEP = 1024.0                    # SORT_ORDER spacing, as USP_EXPT_SHEET_ROWS_OP renumbers (1024, 2048, ...)
MAX_GRID_CELL = 4000                  # ExportSheetController.MaxGridCellChars
MAX_GRID_JSON = 200000                # ExportSheetController.MaxJsonChars (one grid row's DATA)
MAX_META_JSON = 1000000               # ExportSheetController.MaxMetaChars
MAX_LOG_TEXT = 2000                   # TBL_EXPT_SHEET_LOG.NEW_VALUE NVARCHAR(2000)
DAL_AMP = '|;amp;|'                   # DAL.AssignParameterValues rewrites this text; the API guards JSON against it

# Colour CODE legend = the seeds of EXPORT_SHEET_VIEW_LINKSDB20_20261021.sql step 3a (code, label, BG, FG)
STATUSES = [
    ('COMPLETE', 'Shipment Complete', '#ffff00', None),
    ('SI_DONE', 'SI Filling Complete', '#00ff00', None),
    ('BKG_RELEASE', 'Booking Release', '#46bdc6', None),
    ('HAZ', 'HAZ Container (Yes)', None, '#ff0000'),
    ('PLAN_CHANGED', 'Bookings to Planned', '#ff00ff', None),
    ('BKG_AVAILABLE', 'Booking available to be released', '#ff9900', None),
    ('TAKE_BKG', 'Take Booking', '#9fc5e8', None),
    ('LC_AWAITED', 'LC Awaited / Cargo Not Ready', '#ff0000', None),
    ('HOLD', 'HOLD', '#fbbc04', None),
    ('UNKNOWN', 'Unknown', '#9900ff', None),
]
HAZ_YES = ('YES', 'Y', 'HAZ')         # USP_EXPT_SHEET_SAVE_APPLY: HAZ YES / HAZ / Y -> Y (the UI turns those rows red)

# tabs that are never imported (normalised titles)
NEVER_READ = {'column access'}        # staff e-mail addresses: not a single cell of it is read
LEGEND_TITLES = {'colour code', 'color code'}
SKIP_KINDS = {'legend', 'access', 'kpi'}

# expected shape of the target tables (name, column, sys.types name, sys.columns.max_length or None)
def _expected_columns(reg_keys, ml_keys):
    cols = [
        ('TBL_EXPT_SHEET_MONTH', 'ID', 'int', 4), ('TBL_EXPT_SHEET_MONTH', 'CMPCODE', 'varchar', 6),
        ('TBL_EXPT_SHEET_MONTH', 'CITYCODE', 'varchar', 3), ('TBL_EXPT_SHEET_MONTH', 'SHEET_MONTH', 'char', 7),
        ('TBL_EXPT_SHEET_MONTH', 'TITLE', 'nvarchar', 200), ('TBL_EXPT_SHEET_MONTH', 'TABS', 'nvarchar', -1),
        ('TBL_EXPT_SHEET_MONTH', 'ROWVER', 'int', 4), ('TBL_EXPT_SHEET_MONTH', 'MAKER_ID', 'varchar', 50),
        ('TBL_EXPT_SHEET_MONTH', 'MAKER_DT', 'datetime', 8), ('TBL_EXPT_SHEET_MONTH', 'MODIFY_ID', 'varchar', 50),
        ('TBL_EXPT_SHEET_MONTH', 'MODIFY_DT', 'datetime', 8),
        ('TBL_EXPT_SHEET_ROW', 'ID', 'int', 4), ('TBL_EXPT_SHEET_ROW', 'MONTH_ID', 'int', 4),
        ('TBL_EXPT_SHEET_ROW', 'TAB_CODE', 'varchar', 20), ('TBL_EXPT_SHEET_ROW', 'SORT_ORDER', 'float', 8),
        ('TBL_EXPT_SHEET_ROW', 'STATUS_CODE', 'varchar', 20), ('TBL_EXPT_SHEET_ROW', 'EXPTNO', 'varchar', 25),
    ]
    for k in reg_keys:
        cols.append(('TBL_EXPT_SHEET_ROW', k, 'nvarchar', 4000 if k in ml_keys else 1000))
    cols += [
        ('TBL_EXPT_SHEET_ROW', 'CELL_FMT', 'nvarchar', -1), ('TBL_EXPT_SHEET_ROW', 'ROWVER', 'int', 4),
        ('TBL_EXPT_SHEET_ROW', 'ISDELETED', 'bit', 1), ('TBL_EXPT_SHEET_ROW', 'MAKER_ID', 'varchar', 50),
        ('TBL_EXPT_SHEET_ROW', 'MAKER_DT', 'datetime', 8), ('TBL_EXPT_SHEET_ROW', 'MODIFY_ID', 'varchar', 50),
        ('TBL_EXPT_SHEET_ROW', 'MODIFY_DT', 'datetime', 8),
        ('TBL_EXPT_SHEET_GRIDROW', 'ID', 'int', 4), ('TBL_EXPT_SHEET_GRIDROW', 'MONTH_ID', 'int', 4),
        ('TBL_EXPT_SHEET_GRIDROW', 'TAB_CODE', 'varchar', 20), ('TBL_EXPT_SHEET_GRIDROW', 'SORT_ORDER', 'float', 8),
        ('TBL_EXPT_SHEET_GRIDROW', 'DATA', 'nvarchar', -1), ('TBL_EXPT_SHEET_GRIDROW', 'CELL_FMT', 'nvarchar', -1),
        ('TBL_EXPT_SHEET_GRIDROW', 'ROWVER', 'int', 4), ('TBL_EXPT_SHEET_GRIDROW', 'ISDELETED', 'bit', 1),
        ('TBL_EXPT_SHEET_GRIDROW', 'MAKER_ID', 'varchar', 50), ('TBL_EXPT_SHEET_GRIDROW', 'MAKER_DT', 'datetime', 8),
        ('TBL_EXPT_SHEET_GRIDROW', 'MODIFY_ID', 'varchar', 50), ('TBL_EXPT_SHEET_GRIDROW', 'MODIFY_DT', 'datetime', 8),
        ('TBL_EXPT_SHEET_TABMETA', 'MONTH_ID', 'int', 4), ('TBL_EXPT_SHEET_TABMETA', 'TAB_CODE', 'varchar', 20),
        ('TBL_EXPT_SHEET_TABMETA', 'META', 'nvarchar', -1), ('TBL_EXPT_SHEET_TABMETA', 'ROWVER', 'int', 4),
        ('TBL_EXPT_SHEET_TABMETA', 'MODIFY_ID', 'varchar', 50), ('TBL_EXPT_SHEET_TABMETA', 'MODIFY_DT', 'datetime', 8),
        ('TBL_EXPT_SHEET_STATUS', 'CODE', 'varchar', 20), ('TBL_EXPT_SHEET_STATUS', 'BG', 'char', 7),
        ('TBL_EXPT_SHEET_STATUS', 'FG', 'char', 7), ('TBL_EXPT_SHEET_STATUS', 'ACTIVE', 'bit', 1),
        ('TBL_EXPT_SHEET_LOG', 'ID', 'bigint', 8), ('TBL_EXPT_SHEET_LOG', 'ROW_KIND', 'char', 1),
        ('TBL_EXPT_SHEET_LOG', 'ROW_ID', 'int', 4), ('TBL_EXPT_SHEET_LOG', 'COL_KEY', 'varchar', 40),
        ('TBL_EXPT_SHEET_LOG', 'OLD_VALUE', 'nvarchar', 4000), ('TBL_EXPT_SHEET_LOG', 'NEW_VALUE', 'nvarchar', 4000),
        ('TBL_EXPT_SHEET_LOG', 'ACTION', 'varchar', 20), ('TBL_EXPT_SHEET_LOG', 'USER_ID', 'varchar', 50),
        ('TBL_EXPT_SHEET_LOG', 'USER_NAME', 'nvarchar', 200), ('TBL_EXPT_SHEET_LOG', 'LOG_DT', 'datetime', 8),
        ('TBL_EXPT_SHEET_VIEW', 'ID', 'int', 4), ('TBL_EXPT_SHEET_VIEW', 'MONTH_ID', 'int', 4),
    ]
    return cols


class Refused(Exception):
    """The workbook cannot be imported as it is; nothing is written."""


# ======================================================================================================================
# sheet-registry.ts (TypeScript object literals -> Python)
# ======================================================================================================================
def _ts_to_json(src):
    """The subset of TS literal syntax the registry uses -> JSON text: quotes bare keys, turns '..' into "..",
    drops comments and trailing commas."""
    out = []
    i, n = 0, len(src)
    while i < n:
        c = src[i]
        if c in '"\'':
            q, j, buf = c, i + 1, []
            while j < n and src[j] != q:
                if src[j] == '\\' and j + 1 < n:
                    nxt = src[j + 1]
                    if q == "'" and nxt == "'":
                        buf.append("'")
                    else:
                        buf.append(src[j:j + 2])
                    j += 2
                    continue
                buf.append('\\"' if (q == "'" and src[j] == '"') else src[j])
                j += 1
            if j >= n:
                raise ValueError('unterminated string in registry literal')
            out.append('"' + ''.join(buf) + '"')
            i = j + 1
            continue
        if src.startswith('//', i):
            j = src.find('\n', i)
            i = n if j < 0 else j
            continue
        if src.startswith('/*', i):
            j = src.find('*/', i + 2)
            i = n if j < 0 else j + 2
            continue
        if c == ',':
            j = i + 1
            while j < n and src[j] in ' \t\r\n':
                j += 1
            if j < n and src[j] in '}]':
                i = j
                continue
            out.append(c)
            i += 1
            continue
        m = re.match(r'[A-Za-z_$][\w$]*', src[i:])
        if m:
            word = m.group(0)
            k = i + len(word)
            if re.match(r'\s*:', src[k:]) and word not in ('true', 'false', 'null'):
                out.append('"' + word + '"')
            else:
                out.append(word)
            i = k
            continue
        out.append(c)
        i += 1
    return ''.join(out)


def _ts_const(text, name):
    m = re.search(r'export\s+const\s+' + re.escape(name) + r'\b[^=]*=\s*', text)
    if not m:
        raise ValueError('constant %s not found in the registry' % name)
    i = m.end()
    if text[i] not in '[{':
        raise ValueError('constant %s is not an object / array literal' % name)
    depth, j, q = 0, i, None
    while j < len(text):
        c = text[j]
        if q:
            if c == '\\':
                j += 2
                continue
            if c == q:
                q = None
        elif c in '"\'':
            q = c
        elif c in '[{':
            depth += 1
        elif c in ']}':
            depth -= 1
            if depth == 0:
                break
        j += 1
    import json
    return json.loads(_ts_to_json(text[i:j + 1]))


class Registry:
    def __init__(self, path):
        text = open(path, encoding='utf-8').read()
        self.path = path
        self.columns = _ts_const(text, 'SHEET_COLUMNS')
        self.layout_a = _ts_const(text, 'LAYOUT_A')
        self.layout_b = _ts_const(text, 'LAYOUT_B')
        self.default_tabs = _ts_const(text, 'DEFAULT_TABS')
        self.optional_wo_tabs = _ts_const(text, 'OPTIONAL_WO_TABS')
        self.grid_tabs = _ts_const(text, 'GRID_TABS')
        try:
            self.tab_widths = _ts_const(text, 'TAB_WIDTHS')
        except ValueError:
            self.tab_widths = {}
        self.col = {c['key']: c for c in self.columns}
        self.keys = [c['key'] for c in self.columns]
        # DB column sizes (EXPORT_SHEET_VIEW_LINKSDB20_20261021.sql): the multi-line columns are NVARCHAR(2000), the rest 500
        self.ml_keys = {c['key'] for c in self.columns if c['kind'] == 'ml'}
        self.maxlen = {k: (2000 if k in self.ml_keys else 500) for k in self.keys}
        if len(self.keys) != 52 or set(self.layout_a) != set(self.keys):
            raise ValueError('registry: expected the 52 WO columns with LAYOUT_A covering them all')
        if not set(self.layout_b) <= set(self.keys):
            raise ValueError('registry: LAYOUT_B names an unknown key')
        if self.ml_keys != {'VEH', 'REMARK1', 'CONTAINER_NO'}:
            raise ValueError('registry: multi-line columns are %s, the DB has NVARCHAR(2000) for VEH, REMARK1, CONTAINER_NO'
                             % sorted(self.ml_keys))

    def layout(self, name):
        return self.layout_b if name == 'B' else self.layout_a


# ======================================================================================================================
# value rules
# ======================================================================================================================
def sql_key_order(k):
    """ORDER BY on a VARCHAR column under SQL_Latin1_General_CP1_CI_AS (LinksDB20 / tempdb) for [A-Z0-9_] keys:
    '_' sorts before digits, digits before letters, case-insensitive (PO_NO < POL_GATE < PORT)."""
    out = []
    for ch in k.upper():
        if ch == '_':
            out.append(0)
        elif ch.isdigit():
            out.append(1 + ord(ch) - 48)
        elif 'A' <= ch <= 'Z':
            out.append(11 + ord(ch) - 65)
        else:
            out.append(100 + ord(ch))
    return out


def is_sql_date(d, m, y):
    """ISDATE('dd-mm-yyyy') under SET DATEFORMAT dmy (DATETIME range 1753-9999)."""
    if not (1753 <= y <= 9999):
        return False
    try:
        datetime.date(y, m, d)
        return True
    except ValueError:
        return False


RE_DMY = re.compile(r'([0-3][0-9])([-/.])([01][0-9])([-/.])([12][0-9]{3})')
RE_YMD = re.compile(r'([12][0-9]{3})-([01][0-9])-([0-3][0-9])')


def normalize_date_text(v):
    """USP_EXPT_SHEET_SAVE_APPLY on a date column: a real date (dd-mm-yyyy, dd/mm/yyyy, dd.mm.yyyy or yyyy-mm-dd, after
    LTRIM/RTRIM of spaces) becomes dd-mm-yyyy; anything else is kept exactly as typed."""
    s = v.strip(' ')
    m = RE_DMY.fullmatch(s)
    if m and m.group(2) == m.group(4) and is_sql_date(int(m.group(1)), int(m.group(3)), int(m.group(5))):
        return '%s-%s-%s' % (m.group(1), m.group(3), m.group(5))
    m = RE_YMD.fullmatch(s)
    if m and is_sql_date(int(m.group(3)), int(m.group(2)), int(m.group(1))):
        return '%s-%s-%s' % (m.group(3), m.group(2), m.group(1))
    return v


def fmt_date(d):
    return '%02d-%02d-%04d' % (d.day, d.month, d.year)


def num_text(v, warn=None):
    """A number cell as the sheet shows it in General format: 6.0 -> 6, 301102610004227.0 -> 301102610004227, 1.5 -> 1.5."""
    if isinstance(v, bool):
        return 'TRUE' if v else 'FALSE'
    if isinstance(v, int):
        return str(v)
    if v != v or v in (float('inf'), float('-inf')):
        return str(v)
    if v.is_integer():
        if abs(v) >= 1e15 and warn is not None and abs(v) > 2 ** 53:
            warn('number %r is beyond 2^53: the workbook itself may have lost digits' % v)
        return str(int(v))
    r = repr(v)
    if 'e' in r or 'E' in r:
        r = format(decimal.Decimal(r), 'f')
    return r


def xml_safe(s):
    """ExportSheetController.XmlSafe: CR / CRLF folded to LF, characters XML 1.0 cannot carry dropped.
    Returns (text, folded CRs, dropped characters)."""
    folded = s.count('\r')
    if folded:
        s = s.replace('\r\n', '\n').replace('\r', '\n')
    out, dropped = [], 0
    for ch in s:
        o = ord(ch)
        if ch in '\t\n' or 0x20 <= o <= 0xD7FF or 0xE000 <= o <= 0xFFFD or 0x10000 <= o <= 0x10FFFF:
            out.append(ch)
        else:
            dropped += 1
    return ''.join(out), folded, dropped


def nj_str(s):
    """Json.NET string, StringEscapeHandling.Default."""
    out = ['"']
    for ch in s:
        o = ord(ch)
        if ch == '"':
            out.append('\\"')
        elif ch == '\\':
            out.append('\\\\')
        elif ch == '\n':
            out.append('\\n')
        elif ch == '\r':
            out.append('\\r')
        elif ch == '\t':
            out.append('\\t')
        elif ch == '\f':
            out.append('\\f')
        elif ch == '\b':
            out.append('\\b')
        elif o < 0x20 or o in (0x85, 0x2028, 0x2029):
            out.append('\\u%04x' % o)
        else:
            out.append(ch)
    out.append('"')
    return ''.join(out)


def nj(o):
    """JToken.ToString(Formatting.None) of dict (insertion order) / list / str / int / bool / None."""
    if o is None:
        return 'null'
    if isinstance(o, bool):
        return 'true' if o else 'false'
    if isinstance(o, int):
        return str(o)
    if isinstance(o, float):
        if o.is_integer():
            return str(int(o)) + '.0'
        return repr(o)
    if isinstance(o, str):
        return nj_str(o)
    if isinstance(o, dict):
        return '{' + ','.join(nj_str(str(k)) + ':' + nj(v) for k, v in o.items()) + '}'
    if isinstance(o, (list, tuple)):
        return '[' + ','.join(nj(v) for v in o) + ']'
    raise TypeError(type(o))


def guard_json(j):
    """ExportSheetController.GuardJson."""
    return j.replace(DAL_AMP, '\\u007C;amp;\\u007C') if DAL_AMP in j else j


def cell_fmt_json(fmts, order):
    """The one canonical CELL_FMT form USP_EXPT_SHEET_SAVE_FMT writes: {"KEY":{"bg":"#rrggbb","fg":"#rrggbb","b":1},...},
    keys in SQL ORDER BY order, colours lower case, empty cells left out, NULL when nothing is left."""
    parts = []
    for k in sorted(fmts, key=order):
        f = fmts[k]
        frag = []
        if f.get('bg'):
            frag.append('"bg":"%s"' % f['bg'].lower())
        if f.get('fg'):
            frag.append('"fg":"%s"' % f['fg'].lower())
        if f.get('b'):
            frag.append('"b":1')
        if frag:
            parts.append('"%s":{%s}' % (k, ','.join(frag)))
    return '{' + ','.join(parts) + '}' if parts else None


def cell_fmt_is_canonical(j):
    """Port of SAVE_FMT's check of a stored CELL_FMT (else every later format edit of the row is refused as invalid)."""
    if j in (None, '', '{}'):
        return True
    if not (j.startswith('{"') and j.endswith('}}')) or re.search(r'[^A-Za-z0-9_#":,{}]', j):
        return False
    x = j
    for a, b in (('},"', '/><c k="'), ('":{', '" '), ('"bg":"', 'bg="'), ('"fg":"', 'fg="'), ('"b":1', 'b="1"'),
                 ('"b":0', 'b="0"'), (',', ' '), ('{"', '<c k="'), ('}}', '/>')):
        x = x.replace(a, b)
    return (x.startswith('<c k="') and x.endswith('/>') and '{' not in x and '}' not in x and ':' not in x
            and x.count('"') % 2 == 0)


def sql_n(s, chunk=1000):
    """NVARCHAR literal for s, pure ASCII: quotes doubled; control and non-ASCII characters as NCHAR(n) (surrogate
    pairs for characters above U+FFFF), so line-ending or code-page changes to the .sql file cannot alter the data.
    Anything that is not one plain literal starts with CONVERT(NVARCHAR(MAX), ...) so no concatenation is cut at 4,000."""
    if s is None:
        return 'NULL'
    if s == '':
        return "N''"
    parts, buf = [], []

    def flush():
        if buf:
            parts.append("N'" + ''.join(buf) + "'")
            del buf[:]

    count = 0
    for ch in s:
        o = ord(ch)
        if 32 <= o < 127:
            buf.append("''" if ch == "'" else ch)
            count += 1
            if count >= chunk:
                flush()
                count = 0
        else:
            flush()
            count = 0
            if o > 0xFFFF:
                o -= 0x10000
                parts.append('NCHAR(%d)' % (0xD800 + (o >> 10)))
                parts.append('NCHAR(%d)' % (0xDC00 + (o & 0x3FF)))
            else:
                parts.append('NCHAR(%d)' % o)
    flush()
    if len(parts) == 1 and parts[0].startswith("N'"):
        return parts[0]
    return 'CONVERT(NVARCHAR(MAX), ' + parts[0] + ')' + ''.join(' + ' + p for p in parts[1:])


def sql_a(s):
    """VARCHAR literal for an ASCII code / key."""
    if s is None:
        return 'NULL'
    if not all(32 <= ord(c) < 127 for c in s):
        raise ValueError('not ASCII: %r' % s)
    return "'" + s.replace("'", "''") + "'"


def ascii_comment(s):
    """Text for a /* */ comment: ASCII only, no comment delimiters."""
    return ''.join(c if 32 <= ord(c) < 127 and c not in '*/' else '?' for c in str(s))


def ascii_text(s):
    return ''.join(c if 32 <= ord(c) < 127 else '?' for c in str(s))


def px(width_chars):
    """Column width in px, the same formula the registry generator used for TAB_WIDTHS / GRID_TABS."""
    return int(round((width_chars or 8.43) * 7 + 5))


def col_letter(c):
    s = ''
    while c:
        c, r = divmod(c - 1, 26)
        s = chr(65 + r) + s
    return s


def col_index(letters):
    n = 0
    for ch in letters:
        n = n * 26 + (ord(ch.upper()) - 64)
    return n


def norm_header(h):
    return re.sub(r'\s+', ' ', str(h)).strip().upper() if h is not None else ''


def norm_title(t):
    return re.sub(r'\s+', ' ', str(t)).strip().lower()


# ======================================================================================================================
# workbook reading
# ======================================================================================================================
class Colours:
    """openpyxl Color -> '#rrggbb' (theme colours with tint resolved from the workbook theme), None for auto / none."""

    def __init__(self, wb):
        self.theme = ['FFFFFF', '000000', 'FFFFFF', '000000', '4472C4', 'ED7D31', 'A5A5A5', 'FFC000', '5B9BD5', '70AD47',
                      '0563C1', '954F72']
        try:
            root = ET.fromstring(wb.loaded_theme)
            ns = {'a': 'http://schemas.openxmlformats.org/drawingml/2006/main'}
            cs = root.find('.//a:clrScheme', ns)
            vals = {}
            for el in cs:
                tag = el.tag.split('}')[1]
                sub = el[0]
                vals[tag] = (sub.get('val') if sub.tag.endswith('srgbClr') else sub.get('lastClr')) or '000000'
            order = ['lt1', 'dk1', 'lt2', 'dk2', 'accent1', 'accent2', 'accent3', 'accent4', 'accent5', 'accent6',
                     'hlink', 'folHlink']           # theme colour index 0..11 (lt / dk swapped, as Excel does)
            self.theme = [vals.get(k, d) for k, d in zip(order, self.theme)]
        except Exception:
            pass
        try:
            from openpyxl.styles.colors import COLOR_INDEX
            self.indexed = list(COLOR_INDEX)
        except Exception:
            self.indexed = []

    @staticmethod
    def _tint(hexrgb, tint):
        if not tint:
            return hexrgb
        r, g, b = (int(hexrgb[i:i + 2], 16) / 255.0 for i in (0, 2, 4))
        h, l, s = colorsys.rgb_to_hls(r, g, b)
        l = l * (1 + tint) if tint < 0 else l * (1 - tint) + tint
        r, g, b = colorsys.hls_to_rgb(h, max(0.0, min(1.0, l)), s)
        return '%02X%02X%02X' % (int(round(r * 255)), int(round(g * 255)), int(round(b * 255)))

    def resolve(self, c):
        if c is None:
            return None
        try:
            t = c.type
            if t == 'rgb':
                v = c.rgb
                if not isinstance(v, str) or len(v) < 6:
                    return None
                base = v[-6:]
            elif t == 'theme':
                if c.theme is None or not (0 <= c.theme < len(self.theme)):
                    return None
                base = self.theme[c.theme]
            elif t == 'indexed':
                if c.indexed is None or c.indexed >= 64 or c.indexed >= len(self.indexed):
                    return None                     # 64 / 65 = system foreground / background
                base = self.indexed[c.indexed][-6:]
            else:
                return None
            return '#' + self._tint(base.upper(), c.tint or 0).lower()
        except Exception:
            return None


class Report:
    def __init__(self):
        self.lines = []
        self.refusals = []
        self.warnings = []

    def add(self, s=''):
        self.lines.append(s)

    def warn(self, s):
        self.warnings.append(s)

    def refuse(self, s):
        self.refusals.append(s)


def col_dims(ws):
    """Column widths (characters) and hidden columns by index; openpyxl groups ranges under their first letter."""
    widths, hidden = {}, set()
    for key, d in list(ws.column_dimensions.items()):
        try:
            lo = d.min or col_index(key)
            hi = d.max or lo
        except Exception:
            continue
        for c in range(lo, min(hi, 702) + 1):
            if d.width:
                widths[c] = d.width
            if d.hidden:
                hidden.add(c)
    return widths, hidden


def freeze_of(ws):
    fp = ws.freeze_panes
    if not fp:
        return 0, 0
    m = re.fullmatch(r'([A-Z]+)([0-9]+)', str(fp))
    if not m:
        return 0, 0
    return int(m.group(2)) - 1, col_index(m.group(1)) - 1     # frozen rows, frozen cols


class Cell:
    """One cell as the import sees it: text value and formats."""
    __slots__ = ('text', 'bg', 'fg', 'b', 'formula', 'error', 'kind')


class SheetReader:
    def __init__(self, wb_f, wb_v, colours, rep):
        self.wb_f, self.wb_v, self.colours, self.rep = wb_f, wb_v, colours, rep
        self._merges = {}

    def anchor(self, title, r, c):
        """A cell covered by a merge (not its top-left) looks like the top-left one: its fill / font come from there."""
        m = self._merges.get(title)
        if m is None:
            m = {}
            for rng in self.wb_f[title].merged_cells.ranges:
                for rr in range(rng.min_row, rng.max_row + 1):
                    for cc in range(rng.min_col, rng.max_col + 1):
                        if (rr, cc) != (rng.min_row, rng.min_col):
                            m[(rr, cc)] = (rng.min_row, rng.min_col)
            self._merges[title] = m
        return m.get((r, c), (r, c))

    def read(self, title, r, c, date_col, stats):
        ws_f, ws_v = self.wb_f[title], self.wb_v[title]
        cf = ws_f._cells.get((r, c))
        cv = ws_v._cells.get((r, c))
        cs = ws_f._cells.get(self.anchor(title, r, c))     # where the cell's look comes from
        x = Cell()
        x.text = None
        x.formula = False
        x.error = False
        x.bg = x.fg = None
        x.b = False
        x.kind = None
        if cf is not None and (cf.data_type == 'f' or type(cf.value).__name__ in ('ArrayFormula', 'DataTableFormula')):
            x.formula = True
        if cs is not None:
            fill = cs.fill
            if fill is not None and fill.fill_type:
                if fill.fill_type != 'solid':
                    stats['pattern fills read as their foreground colour'] = stats.get('pattern fills read as their foreground colour', 0) + 1
                x.bg = self.colours.resolve(fill.fgColor)
            font = cs.font
            if font is not None:
                x.fg = self.colours.resolve(font.color)
                x.b = bool(font.b)
        v = cv.value if cv is not None else None
        if cv is not None and cv.data_type == 'e':
            x.error = True
        if v is None:
            return x
        if isinstance(v, bool):
            x.text = 'TRUE' if v else 'FALSE'
        elif isinstance(v, datetime.datetime):
            if v.hour or v.minute or v.second:
                x.text = fmt_date(v) + ' ' + v.strftime('%H:%M' if not v.second else '%H:%M:%S')
                stats['date-times kept with their time (as text)'] = stats.get('date-times kept with their time (as text)', 0) + 1
            else:
                x.text = fmt_date(v)
            x.kind = 'date'
        elif isinstance(v, datetime.date):
            x.text = fmt_date(v)
            x.kind = 'date'
        elif isinstance(v, datetime.time):
            x.text = v.strftime('%H:%M' if not v.second else '%H:%M:%S')
            stats['time-only cells (kept as hh:mm text)'] = stats.get('time-only cells (kept as hh:mm text)', 0) + 1
        elif isinstance(v, datetime.timedelta):
            x.text = str(v)
            stats['duration cells (kept as text)'] = stats.get('duration cells (kept as text)', 0) + 1
        elif isinstance(v, (int, float)):
            x.text = num_text(v, self.rep.warn)
            x.kind = 'num'
            nf = cv.number_format if cv is not None else 'General'
            if nf not in ('General', '@', '0'):
                key = 'number formats not applied (value kept as the plain number): ' + nf
                stats[key] = stats.get(key, 0) + 1
        elif isinstance(v, str):
            if v.strip() == '':
                stats['whitespace-only cells left empty'] = stats.get('whitespace-only cells left empty', 0) + 1
                return x
            x.text = v
            if date_col:
                n = normalize_date_text(v)
                if n != v:
                    stats['date text normalised to dd-mm-yyyy'] = stats.get('date text normalised to dd-mm-yyyy', 0) + 1
                    x.text = n
                    x.kind = 'date'
        else:
            x.text = str(v)
            stats['other values kept as text'] = stats.get('other values kept as text', 0) + 1
        if x.formula:
            stats['formula cells (cached value imported, formula dropped)'] = stats.get('formula cells (cached value imported, formula dropped)', 0) + 1
        if x.error:
            stats['error values (#VALUE! ...) imported as text'] = stats.get('error values (#VALUE! ...) imported as text', 0) + 1
        return x


def last_used_row(ws, cols, start):
    last = 0
    for (r, c), cell in ws._cells.items():
        if r >= start and c in cols and cell.value is not None and not (isinstance(cell.value, str) and cell.value.strip() == ''):
            last = max(last, r)
    return last


# ======================================================================================================================
# the import model
# ======================================================================================================================
class Import:
    def __init__(self, args, reg, rep):
        self.args, self.reg, self.rep = args, reg, rep
        self.wo = []            # dicts: seq, tab, sheet_row, sort, status, values{key: text}, fmt json, ff_text, ff_key
        self.wo_tabs = []       # (tabdef, sheet title, rows count)
        self.grid = []          # dicts: seq, tab, sheet_row, sort, data json, fmt json
        self.meta = []          # (tab code, meta json)
        self.tabs = []          # TABS entries
        self.ff = []            # (tab, sheet row, text, key or '')
        self.status_count = {}
        self.unknown_colours = {}
        self.order = sql_key_order

    # ---------------------------------------------------------------------------------------------------------------
    def run(self):
        from openpyxl import load_workbook
        a, rep, reg = self.args, self.rep, self.reg
        with warnings.catch_warnings():
            warnings.simplefilter('ignore')
            wb_f = load_workbook(a.workbook, data_only=False)
            wb_v = load_workbook(a.workbook, data_only=True)
        self.colours = Colours(wb_f)
        self.reader = SheetReader(wb_f, wb_v, self.colours, rep)
        self.status_by_bg = {bg: code for code, _, bg, _ in STATUSES if bg}

        by_title = {}
        for t in reg.default_tabs + reg.optional_wo_tabs:
            by_title[norm_title(t['title'])] = t
        skip_user = {norm_title(s) for s in (a.skip_tab or [])}

        found = {}          # code -> (tabdef, sheet title)
        order_in_wb = []
        rep.add('Workbook tabs:')
        for ws in wb_f.worksheets:
            nt = norm_title(ws.title)
            if nt in NEVER_READ:
                rep.add('  %-30s skipped, never read (staff e-mail addresses)' % ascii_text(ws.title.strip()))
                continue
            if nt in skip_user:
                rep.add('  %-30s skipped (--skip-tab)' % ascii_text(ws.title.strip()))
                continue
            t = by_title.get(nt)
            if t is None:
                if not any(c.value is not None for c in ws._cells.values()):
                    rep.add('  %-30s skipped (empty, not a Sheet view tab)' % ascii_text(ws.title.strip()))
                    continue
                rep.refuse('tab "%s" has data but is not a Sheet view tab (sheet-registry.ts DEFAULT_TABS / OPTIONAL_WO_TABS); '
                           'add it to the registry or pass --skip-tab "%s"' % (ws.title.strip(), ws.title.strip()))
                continue
            if t['code'] in found:
                rep.refuse('two workbook tabs map to %s ("%s" and "%s")' % (t['code'], found[t['code']][1], ws.title))
                continue
            if t['kind'] == 'legend':
                self.check_legend(ws.title)
                rep.add('  %-30s legend: compared with the seeded Colour CODE (not imported)' % ascii_text(ws.title.strip()))
            elif t['kind'] == 'kpi':
                rep.add('  %-30s skipped (computed by the Sheet view: USP_EXPT_SHEET_KPI)' % ascii_text(ws.title.strip()))
            elif t['kind'] == 'access':
                rep.add('  %-30s skipped' % ascii_text(ws.title.strip()))
            found[t['code']] = (t, ws.title)
            order_in_wb.append(t['code'])

        # TABS: DEFAULT_TABS order; WO tabs only when the workbook has them; NIGERIA (optional) right after the WO tab
        # that precedes it in the workbook
        tabs = []
        for t in reg.default_tabs:
            if t['kind'] == 'wo' and t['code'] not in found:
                rep.warn('WO tab %s ("%s") is not in the workbook: left out of the month\'s TABS' % (t['code'], t['title']))
                continue
            tabs.append(t)
        for t in reg.optional_wo_tabs:
            if t['code'] not in found:
                continue
            pos = order_in_wb.index(t['code'])
            prev = [c for c in order_in_wb[:pos] if found[c][0]['kind'] == 'wo']
            idx = None
            if prev:
                idx = next((i for i, x in enumerate(tabs) if x['code'] == prev[-1]), None)
            if idx is None:
                idx = max([i for i, x in enumerate(tabs) if x['kind'] == 'wo'] or [len(tabs) - 1])
            tabs.insert(idx + 1, t)

        rep.add('')
        rep.add('Imported tabs:')
        seq = 0
        for t in tabs:
            entry = {'code': t['code'], 'title': t['title'], 'kind': t['kind']}
            if t['kind'] == 'wo':
                entry['layout'] = t.get('layout') or 'A'
                if t.get('helper'):
                    entry['helper'] = True
                widths, frozen_cols, n = self.wo_tab(t, found[t['code']][1], seq)
                seq += n
                entry['widths'] = widths
                entry['frozenCols'] = frozen_cols
            elif t['kind'] in ('grid', 'free') and t['code'] in found:
                self.grid_tab(t, found[t['code']][1])
            elif t['kind'] in ('grid', 'free'):
                rep.warn('tab %s ("%s") is not in the workbook: it starts empty' % (t['code'], t['title']))
            self.tabs.append(entry)
        self.tabs_json = guard_json(nj(self.tabs))
        if len(self.tabs_json) > MAX_GRID_JSON:
            rep.refuse('TABS JSON is %d characters (API limit %d)' % (len(self.tabs_json), MAX_GRID_JSON))

    # ---------------------------------------------------------------------------------------------------------------
    def check_legend(self, title):
        """Only the legend cells (column A label, column B fill / font) are read; differences are reported."""
        ws = self.reader.wb_f[title]
        seeds = {norm_title(lbl): (code, bg, fg) for code, lbl, bg, fg in STATUSES}
        seen = set()
        for r in range(2, min(ws.max_row, 40) + 1):
            a = ws._cells.get((r, 1))
            if a is None or not isinstance(a.value, str) or not a.value.strip():
                continue
            s = seeds.get(norm_title(a.value))
            if s is None:
                self.rep.warn('Colour CODE row %d "%s" is not one of the 10 seeded statuses' % (r, a.value.strip()))
                continue
            code, bg, fg = s
            seen.add(code)
            b = ws._cells.get((r, 2))
            wbg = self.colours.resolve(b.fill.fgColor) if b is not None and b.fill is not None and b.fill.fill_type else None
            wfg = self.colours.resolve(b.font.color) if b is not None and b.font is not None else None
            if bg and wbg != bg:
                self.rep.warn('Colour CODE %s: workbook colour %s, seeded %s - rows in %s are matched to %s as well'
                              % (code, wbg, bg, wbg, code))
                if wbg and wbg not in ('#ffffff',) and wbg not in self.status_by_bg:
                    self.status_by_bg[wbg] = code
            if fg and wfg != fg and not bg:
                self.rep.warn('Colour CODE %s: workbook font colour %s, seeded %s' % (code, wfg, fg))
        missing = [c for c, _, _, _ in STATUSES if c not in seen]
        if missing:
            self.rep.warn('Colour CODE tab lacks %s (the seeded legend has them)' % ', '.join(missing))

    # ---------------------------------------------------------------------------------------------------------------
    def wo_tab(self, t, title, seq0):
        rep, reg = self.rep, self.reg
        code, layout = t['code'], t.get('layout') or 'A'
        keys = reg.layout(layout)
        ws_f = self.reader.wb_f[title]
        maxc = ws_f.max_column
        stats = {}
        # 1. header row -> registry keys (duplicates by occurrence: 1st "Type" = CNTR_TYPE, 2nd = CARGO_TYPE, ...)
        expected = [(k, norm_header(reg.col[k]['header'])) for k in keys]
        other = {norm_header(reg.col[k]['header']): k for k in reg.keys if k not in keys}
        used, colmap, helpers = set(), {}, []
        data_cols_any = set(c for (r, c), cell in ws_f._cells.items() if r >= 2 and cell.value is not None
                            and not (isinstance(cell.value, str) and not cell.value.strip()))
        if norm_header(ws_f._cells.get((1, 1)).value if ws_f._cells.get((1, 1)) else None) != 'WO':
            rep.refuse('%s: cell A1 is not the WO header - not a WO tab layout' % code)
            return {}, 0, 0
        for c in range(1, maxc + 1):
            hc = ws_f._cells.get((1, c))
            h = hc.value if hc is not None else None
            if hc is not None and (hc.data_type == 'f' or (isinstance(h, str) and h.startswith('='))):
                helpers.append(col_letter(c))
                continue
            nh = norm_header(h)
            if not nh:
                if c in data_cols_any:
                    rep.refuse('%s: column %s has values but no header - cannot map it to a registry column'
                               % (code, col_letter(c)))
                continue
            k = next((k for k, eh in expected if eh == nh and k not in used), None)
            if k:
                colmap[c] = k
                used.add(k)
            elif nh in other:
                msg = '%s: column %s "%s" is %s, which layout %s does not have' % (code, col_letter(c), nh, other[nh], layout)
                if c in data_cols_any:
                    rep.refuse(msg + ' - and it holds values')
                else:
                    rep.warn(msg + ' (empty, skipped)')
            else:
                msg = '%s: column %s header "%s" is not a registry column' % (code, col_letter(c), nh)
                if c in data_cols_any:
                    rep.refuse(msg + ' - and it holds values')
                else:
                    rep.warn(msg + ' (empty, skipped)')
        missing = [k for k in keys if k not in used]
        if missing:
            rep.warn('%s: registry columns not in the workbook (left empty): %s' % (code, ', '.join(missing)))
        if helpers:
            rep.warn('%s: formula helper columns %s not imported (the Sheet view shows the helper list live)'
                     % (code, ', '.join(helpers)))
        # 2. widths / freeze for TABS
        widths_c, hidden_c = col_dims(ws_f)
        widths = {}
        for c, k in sorted(colmap.items(), key=lambda kv: keys.index(kv[1])):
            widths[k] = px(widths_c.get(c))
        hidden_keys = [colmap[c] for c in sorted(hidden_c) if c in colmap]
        if hidden_keys:
            rep.warn('%s: hidden columns not carried (the Sheet view hides columns per user): %s' % (code, ', '.join(hidden_keys)))
        reg_w = reg.tab_widths.get(code)
        if reg_w:
            diff = [k for k in widths if reg_w.get(k) not in (None, widths[k])]
            if diff:
                rep.warn('%s: %d column widths differ from TAB_WIDTHS in sheet-registry.ts, which rounded widths and missed '
                         'grouped columns (the workbook\'s exact widths go into TABS): %s'
                         % (code, len(diff), ', '.join(diff[:8]) + (' ...' if len(diff) > 8 else '')))
        frozen_rows, frozen_cols = freeze_of(ws_f)
        # 3. rows
        mapped = set(colmap)
        last = last_used_row(self.reader.wb_v[title], mapped, 2)
        hidden_rows = sum(1 for r, d in ws_f.row_dimensions.items() if d.hidden and 2 <= r <= last)
        merges = [m for m in ws_f.merged_cells.ranges if m.max_row >= 2 and m.min_row <= max(last, 2)]
        date_keys = {k for k in keys if reg.col[k]['kind'] == 'date'}
        n, empty = 0, 0
        status_tab = {}
        fmt_cells = {'bg': 0, 'fg': 0, 'b': 0}
        for r in range(2, last + 1):
            cells = {k: self.reader.read(title, r, c, k in date_keys, stats) for c, k in colmap.items()}
            values = {k: x.text for k, x in cells.items() if x.text is not None}
            if not values:
                empty += 1
                continue
            n += 1
            # row colour: the fill held by more than half of the row's cells (white = none)
            bgs = {}
            for x in cells.values():
                b = x.bg if x.bg not in (None, '#ffffff') else None
                bgs[b] = bgs.get(b, 0) + 1
            row_bg, cnt = max(bgs.items(), key=lambda kv: kv[1])
            if cnt * 2 <= len(cells):
                row_bg = None
            fgs = {}
            for x in cells.values():
                f = x.fg if x.fg not in (None, '#000000') else None
                fgs[f] = fgs.get(f, 0) + 1
            row_fg, cnt = max(fgs.items(), key=lambda kv: kv[1])
            if cnt * 2 <= len(cells):
                row_fg = None
            haz_yes = (values.get('HAZ') or '').strip().upper() in HAZ_YES
            status = None
            eff_bg = None
            if row_bg is not None:
                status = self.status_by_bg.get(row_bg)
                if status:
                    eff_bg = row_bg
                else:
                    self.unknown_colours.setdefault(row_bg, []).append('%s row %d' % (code, r))
            red_implied = haz_yes
            if status is None and row_bg is None and row_fg == '#ff0000':
                status = 'HAZ'                       # legend: HAZ = red font, no fill
                red_implied = True
            if status:
                self.status_count[status] = self.status_count.get(status, 0) + 1
                status_tab[status] = status_tab.get(status, 0) + 1
            # cell formats that differ from the row (bg on every cell, fg / bold on cells with a value)
            fmts = {}
            for k, x in cells.items():
                f = {}
                cb = x.bg if x.bg != '#ffffff' else None
                if eff_bg is not None:
                    if cb != eff_bg:
                        f['bg'] = cb or '#ffffff'
                elif cb is not None:
                    f['bg'] = cb
                if x.text is not None:
                    fg = x.fg if x.fg not in (None, '#000000') else None
                    if fg and not (fg == '#ff0000' and red_implied):
                        f['fg'] = fg
                    if x.b:
                        f['b'] = 1
                if f:
                    fmts[k] = f
                    for p in f:
                        fmt_cells[p] += 1
            fmt = cell_fmt_json(fmts, self.order)
            if not cell_fmt_is_canonical(fmt):
                raise AssertionError('CELL_FMT not canonical: %r' % fmt)
            # values: XmlSafe (as SaveCells), lengths
            for k in list(values):
                v, folded, dropped = xml_safe(values[k])
                if folded:
                    stats['CR / CRLF folded to LF'] = stats.get('CR / CRLF folded to LF', 0) + folded
                if dropped:
                    stats['characters XML cannot carry dropped (as the API does)'] = stats.get('characters XML cannot carry dropped (as the API does)', 0) + dropped
                if v == '':
                    del values[k]
                    continue
                values[k] = v
                if len(v) > reg.maxlen[k]:
                    rep.refuse('%s row %d %s: %d characters, the column takes %d (nothing is ever cut)'
                               % (code, r, k, len(v), reg.maxlen[k]))
            ff = values.get('FF_JOB')
            ff_key = ''
            if ff is not None:
                s = ff.strip()
                if re.fullmatch(r'[0-9A-Za-z][0-9A-Za-z/_.\-]{0,24}', s):
                    ff_key = s
                self.ff.append((code, r, ff, ff_key, values.get('WO') or ''))
            log = 'imported from workbook tab "%s" row %d' % (title.strip(), r)
            self.wo.append({'seq': seq0 + n, 'tab': code, 'sheet_row': r, 'sort': n * SORT_STEP, 'status': status,
                            'values': values, 'fmt': fmt, 'ff_key': ff_key or None, 'log': log[:400]})
        for m in merges:
            ks = [colmap.get(c) for c in range(m.min_col, m.max_col + 1) if colmap.get(c)]
            if ks:
                rep.warn('%s: merged cells %s (%s) - the value stays on row %d only, the other row(s) of the merge are empty'
                         % (code, m.coord, ', '.join(ks), m.min_row))
        self.wo_tabs.append((t, title, n))
        line = '  %-11s %4d rows' % (code, n)
        if empty:
            line += ', %d empty rows skipped' % empty
        if hidden_rows:
            line += ', %d hidden rows (imported; the Sheet view does not hide rows)' % hidden_rows
        rep.add(line)
        if status_tab:
            rep.add('              statuses: ' + ', '.join('%s %d' % kv for kv in sorted(status_tab.items())))
        rep.add('              cell formats kept: %d fills, %d font colours, %d bold' % (fmt_cells['bg'], fmt_cells['fg'], fmt_cells['b']))
        for k in sorted(stats):
            rep.add('              %s: %d' % (k, stats[k]))
        return widths, frozen_cols, n

    # ---------------------------------------------------------------------------------------------------------------
    def grid_tab(self, t, title):
        rep, reg = self.rep, self.reg
        code = t['code']
        gdef = reg.grid_tabs.get(code)
        if gdef is None:
            rep.refuse('%s: no GRID_TABS entry in sheet-registry.ts' % code)
            return
        header_row = bool(gdef.get('headerRow'))
        kinds = {c['key']: c.get('kind') for c in gdef.get('cols', [])}
        reg_headers = {c['key']: c.get('header') or '' for c in gdef.get('cols', [])}
        ws_f, ws_v = self.reader.wb_f[title], self.reader.wb_v[title]
        stats = {}
        used_cols = set(c for (r, c), cell in ws_v._cells.items()
                        if cell.value is not None and not (isinstance(cell.value, str) and not cell.value.strip()))
        widths_c, hidden_c = col_dims(ws_f)
        merges = list(ws_f.merged_cells.ranges)
        last_col = max(list(used_cols) + [m.max_col for m in merges] + [1])      # widths beyond the data are not columns
        start = 2 if header_row else 1
        last = last_used_row(ws_v, set(range(1, last_col + 1)), start)
        frozen_rows, frozen_cols = freeze_of(ws_f)
        meta = {'headerRow': header_row, 'frozenRows': frozen_rows, 'frozenCols': frozen_cols}
        if header_row:
            hdr = {}
            for c in range(1, last_col + 1):
                L = col_letter(c)
                hc = ws_v._cells.get((1, c))
                h = re.sub(r'\s+', ' ', str(hc.value)).strip() if hc is not None and hc.value is not None else ''
                rh = reg_headers.get(L, '')
                if h:
                    hdr[L] = h
                    if rh and norm_header(rh) != norm_header(h):
                        rep.warn('%s: header %s is "%s" in the workbook, "%s" in GRID_TABS (the workbook\'s is kept)' % (code, L, h, rh))
                elif rh:
                    hdr[L] = rh
            meta['headers'] = hdr
        meta['widths'] = {col_letter(c): px(widths_c.get(c)) for c in range(1, last_col + 1)}
        hid = [col_letter(c) for c in sorted(hidden_c) if c <= last_col]
        if hid:
            meta['hidden'] = hid
        mg = []
        for m in sorted(merges, key=lambda m: (m.min_row, m.min_col)):
            if header_row and m.min_row <= 1:
                rep.warn('%s: merge %s touches the header row - not carried' % (code, m.coord))
                continue
            off = 1 if header_row else 0
            mg.append('%s%d:%s%d' % (col_letter(m.min_col), m.min_row - off, col_letter(m.max_col), m.max_row - off))
        if mg:
            meta['merges'] = mg
        meta_json = guard_json(nj(meta))
        if len(meta_json) > MAX_META_JSON:
            rep.refuse('%s: tab layout JSON is %d characters (API limit %d)' % (code, len(meta_json), MAX_META_JSON))
        self.meta.append((code, meta_json))
        n = empty = 0
        fmt_cells = {'bg': 0, 'fg': 0, 'b': 0}
        not_carried = {}
        for r in range(start, last + 1):
            n += 1
            data, fmts = {}, {}
            for c in range(1, last_col + 1):
                L = col_letter(c)
                x = self.reader.read(title, r, c, kinds.get(L) == 'date', stats)
                if x.text is not None:
                    if len(x.text) > MAX_GRID_CELL:
                        rep.refuse('%s row %d column %s: %d characters, the API takes %d per grid cell'
                                   % (code, r, L, len(x.text), MAX_GRID_CELL))
                    data[L] = x.text
                f = {}
                if x.bg and x.bg != '#ffffff':
                    f['bg'] = x.bg
                if x.text is not None:
                    if x.fg and x.fg != '#000000':
                        f['fg'] = x.fg
                    if x.b:
                        f['b'] = 1
                    cf = ws_f._cells.get((r, c))
                    if cf is not None and cf.alignment is not None and (cf.alignment.horizontal or cf.alignment.wrap_text):
                        not_carried['alignment / wrap'] = not_carried.get('alignment / wrap', 0) + 1
                if f:
                    fmts[L] = f
                    for p in f:
                        fmt_cells[p] += 1
            if not data:
                empty += 1
            dj = guard_json(nj(data)) if data else None
            if dj is not None and len(dj) > MAX_GRID_JSON:
                rep.refuse('%s row %d: row JSON is %d characters (API limit %d)' % (code, r, len(dj), MAX_GRID_JSON))
            fj = cell_fmt_json(fmts, self.order)
            if not cell_fmt_is_canonical(fj):
                raise AssertionError('grid CELL_FMT not canonical: %r' % fj)
            log = 'imported from workbook tab "%s" row %d' % (title.strip(), r)
            self.grid.append({'tab': code, 'sheet_row': r, 'sort': n * SORT_STEP, 'data': dj, 'fmt': fj, 'log': log[:400]})
        hidden_rows = sum(1 for rr, d in ws_f.row_dimensions.items() if d.hidden and start <= rr <= last)
        heights = sum(1 for rr, d in ws_f.row_dimensions.items() if d.height and start <= rr <= last)
        line = '  %-11s %4d grid rows (sheet rows %d-%d)' % (code, n, start, last)
        if empty:
            line += ', %d of them empty (kept: the tab keeps its spacing)' % empty
        rep.add(line)
        rep.add('              layout: %d columns, frozen rows %d, %d merges, %d hidden columns%s'
                % (last_col, frozen_rows, len(mg), len(hid), ', header row' if header_row else ''))
        rep.add('              cell formats kept: %d fills, %d font colours, %d bold' % (fmt_cells['bg'], fmt_cells['fg'], fmt_cells['b']))
        if hidden_rows:
            rep.add('              hidden rows (imported, not hidden in the Sheet view): %d' % hidden_rows)
        if heights:
            not_carried['row heights'] = heights
        for k in sorted(not_carried):
            rep.add('              not carried (no place in CELL_FMT / META): %s on %d cells / rows' % (k, not_carried[k]))
        for k in sorted(stats):
            rep.add('              %s: %d' % (k, stats[k]))


# ======================================================================================================================
# SQL
# ======================================================================================================================
class Sql:
    def __init__(self, imp, args, names, wb_hash):
        self.imp, self.a, self.names, self.wb_hash = imp, args, names, wb_hash
        self.reg = imp.reg
        self.cmp, self.city, self.ym = args.cmpcode, args.citycode, args.month

    # -- shared pieces --------------------------------------------------------------------------------------------
    def structure_checks(self, lines, need_tables=True):
        cols = _expected_columns(self.reg.keys, self.reg.ml_keys)
        lines.append("-- 1a. the database, the switches batch, and every table / column this script uses, with its type and size")
        lines.append("IF UPPER(DB_NAME()) <> 'LINKSDB20'")
        lines.append("    INSERT INTO @problems (problem) VALUES (N'Connected to ' + DB_NAME() + N', not LinksDB20.');")
        lines.append("INSERT INTO @problems (problem)")
        lines.append("SELECT N'dbo.' + x.t + N'.' + x.c + N' is missing or not ' + x.ty")
        lines.append("       + CASE WHEN x.ln = -1 THEN N'(MAX)' WHEN x.ty IN ('nvarchar', 'nchar') THEN N'(' + CONVERT(NVARCHAR(10), x.ln / 2) + N')'")
        lines.append("              WHEN x.ty IN ('varchar', 'char') THEN N'(' + CONVERT(NVARCHAR(10), x.ln) + N')' ELSE N'' END")
        lines.append("       + N' - apply EXPORT_SHEET_VIEW_LINKSDB20_20261021.sql first (or the table has another shape).'")
        lines.append("  FROM (VALUES")
        vals = ["        ('%s', '%s', '%s', %d)" % c for c in cols]
        lines.append(',\n'.join(vals))
        lines.append("       ) x (t, c, ty, ln)")
        lines.append(" WHERE NOT EXISTS (SELECT 1 FROM sys.columns sc JOIN sys.types st ON st.user_type_id = sc.user_type_id")
        lines.append("                    WHERE sc.object_id = OBJECT_ID('dbo.' + x.t, 'U') AND sc.name = x.c AND st.name = x.ty AND sc.max_length = x.ln);")
        lines.append("INSERT INTO @problems (problem)")
        lines.append("SELECT N'Column dbo.' + x.t + N'.' + x.c + N' not found.'")
        lines.append("  FROM (VALUES ('EXPT_MASTER', 'EXPTNO'), ('EXPT_MASTER', 'CMPCODE'), ('EXPT_MASTER', 'CITYCODE'), ('EXPT_MASTER', 'EXPTCANCEL'),")
        lines.append("               ('EXPT_MASTER', 'EXPT_INVNO'),")
        lines.append("               ('CMP_ADDRESS', 'CMP_CODE'), ('CMP_ADDRESS', 'CMP_CITY')) x (t, c)")
        lines.append(" WHERE COL_LENGTH('dbo.' + x.t, x.c) IS NULL;")

    def pristine_checks(self, lines, mid, label):
        """Problems when anything in month @mid was written by someone other than this import (same rule in the
        import's @Replace path and in the rollback)."""
        u = sql_a(IMPORT_USER)
        lines.append("    -- everything in the month must still be this import's own: no one created, edited, formatted, moved, linked or")
        lines.append("    -- deleted anything there since, and no personal view points at it")
        lines.append("    INSERT INTO @problems (problem)")
        lines.append("    SELECT N'%s: the month row was created or changed by ' + ISNULL(M.MAKER_ID, N'?') + N' / ' + ISNULL(M.MODIFY_ID, N'?') + N'.'" % label)
        lines.append("      FROM dbo.TBL_EXPT_SHEET_MONTH M WHERE M.ID = %s AND (ISNULL(M.MAKER_ID, '') <> %s OR ISNULL(M.MODIFY_ID, '') <> %s);" % (mid, u, u))
        for tbl, what in (('TBL_EXPT_SHEET_ROW', 'WO row(s)'), ('TBL_EXPT_SHEET_GRIDROW', 'grid row(s)')):
            lines.append("    INSERT INTO @problems (problem)")
            lines.append("    SELECT N'%s: ' + CONVERT(NVARCHAR(10), COUNT(*)) + N' %s added or changed since the import by: '" % (label, what))
            lines.append("           + ISNULL(STUFF((SELECT DISTINCT N', ' + ISNULL(X.MODIFY_ID, N'?') FROM dbo.%s X" % tbl)
            lines.append("                           WHERE X.MONTH_ID = %s AND (ISNULL(X.MAKER_ID, '') <> %s OR ISNULL(X.MODIFY_ID, '') <> %s)" % (mid, u, u))
            lines.append("                           FOR XML PATH(''), TYPE).value('.', 'NVARCHAR(MAX)'), 1, 2, N''), N'') + N'.'")
            lines.append("      FROM dbo.%s R WHERE R.MONTH_ID = %s AND (ISNULL(R.MAKER_ID, '') <> %s OR ISNULL(R.MODIFY_ID, '') <> %s)" % (tbl, mid, u, u))
            lines.append("    HAVING COUNT(*) > 0;")
        lines.append("    INSERT INTO @problems (problem)")
        lines.append("    SELECT N'%s: ' + CONVERT(NVARCHAR(10), COUNT(*)) + N' tab layout(s) saved since the import.'" % label)
        lines.append("      FROM dbo.TBL_EXPT_SHEET_TABMETA T WHERE T.MONTH_ID = %s AND ISNULL(T.MODIFY_ID, '') <> %s HAVING COUNT(*) > 0;" % (mid, u))
        lines.append("    INSERT INTO @problems (problem)")
        lines.append("    SELECT N'%s: ' + CONVERT(NVARCHAR(10), COUNT(*)) + N' history entr(ies) by other users ('" % label)
        lines.append("           + ISNULL(STUFF((SELECT DISTINCT N', ' + ISNULL(L2.USER_ID, N'?') FROM dbo.TBL_EXPT_SHEET_LOG L2")
        lines.append("                           WHERE ISNULL(L2.USER_ID, '') <> %s AND ((L2.ROW_KIND = 'M' AND L2.ROW_ID = %s)" % (u, mid))
        lines.append("                              OR (L2.ROW_KIND = 'W' AND L2.ROW_ID IN (SELECT ID FROM dbo.TBL_EXPT_SHEET_ROW WHERE MONTH_ID = %s))" % mid)
        lines.append("                              OR (L2.ROW_KIND = 'G' AND L2.ROW_ID IN (SELECT ID FROM dbo.TBL_EXPT_SHEET_GRIDROW WHERE MONTH_ID = %s)))" % mid)
        lines.append("                           FOR XML PATH(''), TYPE).value('.', 'NVARCHAR(MAX)'), 1, 2, N''), N'') + N').'")
        lines.append("      FROM dbo.TBL_EXPT_SHEET_LOG L")
        lines.append("     WHERE ISNULL(L.USER_ID, '') <> %s AND ((L.ROW_KIND = 'M' AND L.ROW_ID = %s)" % (u, mid))
        lines.append("        OR (L.ROW_KIND = 'W' AND L.ROW_ID IN (SELECT ID FROM dbo.TBL_EXPT_SHEET_ROW WHERE MONTH_ID = %s))" % mid)
        lines.append("        OR (L.ROW_KIND = 'G' AND L.ROW_ID IN (SELECT ID FROM dbo.TBL_EXPT_SHEET_GRIDROW WHERE MONTH_ID = %s)))" % mid)
        lines.append("    HAVING COUNT(*) > 0;")
        lines.append("    INSERT INTO @problems (problem)")
        lines.append("    SELECT N'%s: ' + CONVERT(NVARCHAR(10), COUNT(*)) + N' personal filter / layout view(s) point at the month (TBL_EXPT_SHEET_VIEW).'" % label)
        lines.append("      FROM dbo.TBL_EXPT_SHEET_VIEW V WHERE V.MONTH_ID = %s HAVING COUNT(*) > 0;" % mid)

    def match_case(self, key_expr, cmp_expr, city_expr):
        return ("CASE WHEN ISNULL(%(k)s, '') = '' THEN 'not a job number'\n"
                "                 WHEN EXISTS (SELECT 1 FROM dbo.EXPT_MASTER E WHERE E.EXPTNO = %(k)s AND LTRIM(RTRIM(ISNULL(E.CMPCODE, ''))) = %(c)s\n"
                "                              AND LTRIM(RTRIM(ISNULL(E.CITYCODE, ''))) = %(t)s AND ISNULL(E.EXPTCANCEL, '') <> 'Y') THEN 'match'\n"
                "                 WHEN EXISTS (SELECT 1 FROM dbo.EXPT_MASTER E WHERE E.EXPTNO = %(k)s AND LTRIM(RTRIM(ISNULL(E.CMPCODE, ''))) = %(c)s\n"
                "                              AND LTRIM(RTRIM(ISNULL(E.CITYCODE, ''))) = %(t)s) THEN 'cancelled job'\n"
                "                 WHEN EXISTS (SELECT 1 FROM dbo.EXPT_MASTER E WHERE E.EXPTNO = %(k)s)\n"
                "                      THEN 'other company / branch: ' + (SELECT TOP 1 LTRIM(RTRIM(ISNULL(E.CMPCODE, ''))) + '/' + LTRIM(RTRIM(ISNULL(E.CITYCODE, '')))\n"
                "                                                       FROM dbo.EXPT_MASTER E WHERE E.EXPTNO = %(k)s)\n"
                "                 ELSE 'not in EXPT_MASTER' END" % {'k': key_expr, 'c': cmp_expr, 't': city_expr})

    def ff_values(self):
        rows = []
        for tab, r, text, key, wo in self.imp.ff:
            rows.append("('%s', %d, %s, %s, %s)" % (tab, r, sql_n(text), sql_a(key), sql_n(wo)))
        return rows

    def header(self, lines, name, why, body):
        lines.append('/*' + '=' * 95)
        lines.append('  %s%s(LinksDB20)' % (name, ' ' * max(9, 58 - len(name))))
        lines.append('')
        for l in why:
            lines.append('  ' + l if l else '')
        lines.append('')
        for l in body:
            lines.append('  ' + l if l else '')
        lines.append('=' * 95 + '*/')

    # -- the import -------------------------------------------------------------------------------------------------
    def import_script(self):
        imp, a, reg = self.imp, self.a, self.reg
        L = []
        per_tab = [(t['code'], n) for t, _, n in imp.wo_tabs]
        grid_tab = {}
        for g in imp.grid:
            grid_tab[g['tab']] = grid_tab.get(g['tab'], 0) + 1
        n_wo, n_g, n_t = len(imp.wo), len(imp.grid), len(imp.meta)
        n_ff = sum(1 for x in imp.wo if x['ff_key'])
        title = ascii_comment(a.title)
        self.header(L, self.names['import'], [
            'Why  : One-time cut-over of the Bajaj monthly Google workbook "%s" (month %s) into Export > Sheet view' % (title, self.ym),
            '       (docs/SHEET_VIEW_SPEC.md section 7 step 6). Generated by links-22/tools/sheet-import/sheet_import.py %s' % VERSION,
            '       on %s from %s' % (datetime.datetime.now().strftime('%Y-%m-%d %H:%M'), ascii_comment(os.path.basename(a.workbook))),
            '       (SHA-256 %s).' % self.wb_hash,
            '       It holds customer shipment data: keep it out of git and delete it after the cut-over.',
        ], [
            'Writes (new rows only - nothing that exists is changed):',
            '  TBL_EXPT_SHEET_MONTH    1 month: CMPCODE %s / CITYCODE %s / %s, TITLE, TABS (the workbook\'s tabs, widths, frozen columns)' % (self.cmp, self.city, self.ym),
            '  TBL_EXPT_SHEET_ROW      %d WO rows in workbook order (SORT_ORDER 1024, 2048, ...): %s' % (n_wo, ', '.join('%s %d' % x for x in per_tab)),
            '                          STATUS_CODE from the row fill (Colour CODE legend), cell fills / font colours / bold in CELL_FMT,',
            '                          dates dd-mm-yyyy (anything else as typed), EXPTNO from FF JOB (%d rows carry a job-like FF JOB)' % n_ff,
            '  TBL_EXPT_SHEET_GRIDROW  %d rows: %s (DATA = Json.NET JSON by column letter)' % (n_g, ', '.join('%s %d' % kv for kv in grid_tab.items())),
            '  TBL_EXPT_SHEET_TABMETA  %d tab layouts (headers, widths, hidden columns, merges, frozen rows / columns)' % n_t,
            '  TBL_EXPT_SHEET_LOG      one \'insert\' entry per month, WO row, grid row and tab layout',
            '  Everything as MAKER_ID / MODIFY_ID / USER_ID / USER_NAME \'%s\'.' % IMPORT_USER,
            '',
            'Switches (step 0 below):',
            '  @Replace        0 = stop when %s already exists for %s / %s;' % (self.ym, self.cmp, self.city),
            '                  1 = replace it, only when it is an earlier run of THIS import that nobody has touched since.',
            '  @MirrorTracking 0 = FF JOB -> EXPTNO only (EXPT_JOB_TRACKING untouched);',
            '                  1 = link every matched job through dbo.USP_EXPT_SHEET_LINK, which also writes the rows\' tracking',
            '                      columns into EXPT_JOB_TRACKING / EXPT_JOB_TRACKING_LOG (not undone by the rollback).',
            '',
            'FF JOB: matched at run time to EXPT_MASTER.EXPTNO of the same company / branch (job not cancelled). Unmatched',
            '        values stay as FF JOB text and are listed in step 1b and in the verify step.',
            '',
            'Safety : step 1 only reads and stops everything (SET NOEXEC ON) unless: the database is LinksDB20, the switches',
            '         batch ran, every sheet table / column exists with the type and size the values were checked against,',
            '         the 10 Colour CODE statuses are seeded, the company / branch exists, the month is free (or @Replace = 1',
            '         and it is an untouched earlier import), and the login may INSERT (and DELETE when replacing) / EXECUTE.',
            '         Step 2 only fills #temp tables; step 3 checks them against the expected counts and writes everything in',
            '         ONE transaction (any error: all of it is rolled back). Step 4 only reads.',
            '         Values were checked by the generator against the column sizes (nothing is cut); text is written as',
            '         N\'\' literals with NCHAR() for line breaks and non-ASCII characters, so the file is plain ASCII.',
            'Rollback: %s (refuses once anyone has changed the imported month)' % self.names['rollback'],
            'Run as : a login with INSERT on the sheet tables (normal login). SQL Server 2008 R2 (no JSON functions, TRY_CONVERT,',
            '         IIF, CONCAT, THROW, inline DECLARE initialisation).',
        ])
        L.append('SET NOEXEC OFF;')
        L.append('GO')
        L.append('USE [LinksDB20]')
        L.append('GO')
        L.append('SET NOCOUNT ON;')
        L.append('SET ANSI_NULLS ON;')
        L.append('SET QUOTED_IDENTIFIER ON;')
        L.append('SET ANSI_PADDING ON;')
        L.append('SET ANSI_WARNINGS ON;')
        L.append('SET ARITHABORT ON;')
        L.append('SET CONCAT_NULL_YIELDS_NULL ON;')
        L.append('GO')
        L.append('')
        L.append('/* ---- 0. switches: edit here ----------------------------------------------------------------- */')
        L.append("IF OBJECT_ID('tempdb..#IMP_OPT') IS NOT NULL DROP TABLE #IMP_OPT;")
        L.append('CREATE TABLE #IMP_OPT (REPLACE_MONTH BIT NOT NULL, MIRROR_TRACKING BIT NOT NULL, PREREQ_OK BIT NOT NULL);')
        L.append('DECLARE @Replace BIT, @MirrorTracking BIT;')
        L.append('SET @Replace        = 0;     -- 1 = replace an earlier, untouched run of this import for the same month')
        L.append('SET @MirrorTracking = 0;     -- 1 = link matched jobs through dbo.USP_EXPT_SHEET_LINK (also writes EXPT_JOB_TRACKING)')
        L.append('INSERT INTO #IMP_OPT (REPLACE_MONTH, MIRROR_TRACKING, PREREQ_OK) VALUES (ISNULL(@Replace, 0), ISNULL(@MirrorTracking, 0), 0);')
        L.append("PRINT 'Switches: @Replace = ' + CONVERT(VARCHAR(1), ISNULL(@Replace, 0)) + ', @MirrorTracking = ' + CONVERT(VARCHAR(1), ISNULL(@MirrorTracking, 0)) + '.';")
        L.append('GO')
        L.append('')
        # 1a structure
        L.append('/* ---- 1. prerequisites (read-only) ---------------------------------------------------------- */')
        L.append('DECLARE @problems TABLE (n INT IDENTITY(1,1), problem NVARCHAR(1000));')
        L.append("IF OBJECT_ID('tempdb..#IMP_OPT') IS NULL")
        L.append("    INSERT INTO @problems (problem) VALUES (N'Step 0 (switches) did not run - run the whole script from the top.');")
        self.structure_checks(L)
        L.append('IF EXISTS (SELECT 1 FROM @problems)')
        L.append('BEGIN')
        L.append("    PRINT '*** Prerequisite check failed on ' + DB_NAME() + ' - NOTHING was written:';")
        L.append('    SELECT n AS [#], problem FROM @problems ORDER BY n;')
        L.append("    RAISERROR('Prerequisite check failed - see the list above.', 16, 1);")
        L.append('    SET NOEXEC ON;     -- skip the rest of the script')
        L.append('END')
        L.append('GO')
        L.append('')
        # 1b data checks
        used_status = sorted({x['status'] for x in imp.wo if x['status']})
        L.append('-- 1b. legend, company / branch, month, permissions, FF JOB preview')
        L.append('DECLARE @problems TABLE (n INT IDENTITY(1,1), problem NVARCHAR(1000));')
        L.append('DECLARE @CMP VARCHAR(6), @CITY VARCHAR(3), @YM CHAR(7), @REPLACE BIT, @MIRROR BIT, @OLD INT;')
        L.append('SELECT @CMP = %s, @CITY = %s, @YM = %s;' % (sql_a(self.cmp), sql_a(self.city), sql_a(self.ym)))
        L.append('SELECT @REPLACE = REPLACE_MONTH, @MIRROR = MIRROR_TRACKING FROM #IMP_OPT;')
        L.append('INSERT INTO @problems (problem)')
        L.append("SELECT N'Colour CODE status ' + x.c + N' is not in dbo.TBL_EXPT_SHEET_STATUS (seeded by EXPORT_SHEET_VIEW_LINKSDB20_20261021.sql).'")
        L.append('  FROM (VALUES %s) x (c)' % ', '.join("('%s')" % s[0] for s in STATUSES))
        L.append(' WHERE NOT EXISTS (SELECT 1 FROM dbo.TBL_EXPT_SHEET_STATUS S WHERE S.CODE = x.c);')
        L.append('-- the legend colours the rows were matched with (a changed colour only re-paints; the status stays right)')
        L.append("SELECT S.CODE, S.BG AS LIVE_BG, x.bg AS WORKBOOK_BG, S.FG AS LIVE_FG, x.fg AS WORKBOOK_FG,")
        L.append("       CASE WHEN ISNULL(S.BG, '') = ISNULL(x.bg, '') AND ISNULL(S.FG, '') = ISNULL(x.fg, '') THEN 'same' ELSE 'legend changed since the seed' END AS NOTE")
        L.append('  FROM (VALUES %s) x (c, bg, fg)' % ', '.join("('%s', %s, %s)" % (c, sql_a(bg) if bg else 'NULL', sql_a(fg) if fg else 'NULL') for c, _, bg, fg in STATUSES))
        L.append('  JOIN dbo.TBL_EXPT_SHEET_STATUS S ON S.CODE = x.c')
        L.append(" WHERE ISNULL(S.BG, '') <> ISNULL(x.bg, '') OR ISNULL(S.FG, '') <> ISNULL(x.fg, '');")
        if used_status:
            L.append('-- statuses the import writes: %s' % ', '.join(used_status))
        L.append('IF NOT EXISTS (SELECT 1 FROM dbo.CMP_ADDRESS WHERE LTRIM(RTRIM(CMP_CODE)) = @CMP AND LTRIM(RTRIM(CMP_CITY)) = @CITY)')
        L.append("    INSERT INTO @problems (problem) VALUES (N'Company ' + @CMP + N' / branch ' + @CITY + N' is not in dbo.CMP_ADDRESS.');")
        L.append('SELECT @OLD = MIN(ID) FROM dbo.TBL_EXPT_SHEET_MONTH WHERE CMPCODE = @CMP AND CITYCODE = @CITY AND SHEET_MONTH = @YM;')
        L.append('IF @OLD IS NOT NULL AND @REPLACE = 0')
        L.append("    INSERT INTO @problems (problem) VALUES (N'Month ' + @YM + N' already exists for ' + @CMP + N' / ' + @CITY + N' (ID ' + CONVERT(NVARCHAR(10), @OLD)")
        L.append("        + N'). Set @Replace = 1 to replace an earlier run of this import, or remove that month first.');")
        L.append('IF @OLD IS NOT NULL AND @REPLACE = 1')
        L.append('BEGIN')
        self.pristine_checks(L, '@OLD', 'Cannot replace month ' + self.ym)
        L.append('END')
        L.append('-- permissions')
        L.append('INSERT INTO @problems (problem)')
        L.append("SELECT N'No ' + x.p + N' permission on dbo.' + x.t + N'.'")
        L.append("  FROM (VALUES ('TBL_EXPT_SHEET_MONTH', 'INSERT', 0), ('TBL_EXPT_SHEET_ROW', 'INSERT', 0), ('TBL_EXPT_SHEET_GRIDROW', 'INSERT', 0),")
        L.append("               ('TBL_EXPT_SHEET_TABMETA', 'INSERT', 0), ('TBL_EXPT_SHEET_LOG', 'INSERT', 0), ('TBL_EXPT_SHEET_STATUS', 'SELECT', 0),")
        L.append("               ('EXPT_MASTER', 'SELECT', 0), ('CMP_ADDRESS', 'SELECT', 0),")
        L.append("               ('TBL_EXPT_SHEET_MONTH', 'DELETE', 1), ('TBL_EXPT_SHEET_ROW', 'DELETE', 1), ('TBL_EXPT_SHEET_GRIDROW', 'DELETE', 1),")
        L.append("               ('TBL_EXPT_SHEET_TABMETA', 'DELETE', 1), ('TBL_EXPT_SHEET_LOG', 'DELETE', 1)) x (t, p, onreplace)")
        L.append(" WHERE (x.onreplace = 0 OR (@REPLACE = 1 AND @OLD IS NOT NULL)) AND ISNULL(HAS_PERMS_BY_NAME('dbo.' + x.t, 'OBJECT', x.p), 0) = 0;")
        L.append('IF @MIRROR = 1')
        L.append('BEGIN')
        L.append("    IF OBJECT_ID('dbo.USP_EXPT_SHEET_LINK', 'P') IS NULL")
        L.append("        INSERT INTO @problems (problem) VALUES (N'@MirrorTracking = 1 needs dbo.USP_EXPT_SHEET_LINK (EXPORT_SHEET_VIEW_LINKSDB20_20261021.sql).');")
        L.append("    ELSE IF ISNULL(HAS_PERMS_BY_NAME('dbo.USP_EXPT_SHEET_LINK', 'OBJECT', 'EXECUTE'), 0) = 0")
        L.append("        INSERT INTO @problems (problem) VALUES (N'No EXECUTE permission on dbo.USP_EXPT_SHEET_LINK (@MirrorTracking = 1).');")
        L.append("    ELSE IF CHARINDEX('EXPORT_SHEET_VIEW_LINKSDB20_20261021', ISNULL(OBJECT_DEFINITION(OBJECT_ID('dbo.USP_EXPT_SHEET_LINK')), '')) = 0")
        L.append("        INSERT INTO @problems (problem) VALUES (N'dbo.USP_EXPT_SHEET_LINK was not created by EXPORT_SHEET_VIEW_LINKSDB20_20261021.sql.');")
        L.append('END')
        ffv = self.ff_values()
        if ffv:
            L.append('-- FF JOB preview: what step 3 will link (read-only)')
            L.append(';WITH V (TAB_CODE, SHEET_ROW, FF_JOB, FF_KEY, WO) AS (SELECT * FROM (VALUES')
            L.append(',\n'.join('        ' + v for v in ffv))
            L.append('    ) X (TAB_CODE, SHEET_ROW, FF_JOB, FF_KEY, WO)),')
            L.append('     R AS (SELECT V.*, ' + self.match_case('V.FF_KEY', '@CMP', '@CITY') + ' AS RESULT FROM V)')
            L.append("SELECT RESULT, COUNT(*) AS SHEET_ROWS, COUNT(DISTINCT NULLIF(FF_KEY, '')) AS DISTINCT_JOBS FROM R GROUP BY RESULT ORDER BY RESULT;")
            L.append("-- rows that will link although their WO is not the job's invoice no. (worth a look before the cut-over)")
            L.append(';WITH V (TAB_CODE, SHEET_ROW, FF_JOB, FF_KEY, WO) AS (SELECT * FROM (VALUES')
            L.append(',\n'.join('        ' + v for v in ffv))
            L.append('    ) X (TAB_CODE, SHEET_ROW, FF_JOB, FF_KEY, WO))')
            L.append(self.wo_check_select('@CMP', '@CITY') + ';')
        L.append('IF EXISTS (SELECT 1 FROM @problems)')
        L.append('BEGIN')
        L.append("    PRINT '*** Prerequisite check failed on ' + DB_NAME() + ' - NOTHING was written:';")
        L.append('    SELECT n AS [#], problem FROM @problems ORDER BY n;')
        L.append("    RAISERROR('Prerequisite check failed - see the list above.', 16, 1);")
        L.append('    SET NOEXEC ON;')
        L.append('END')
        L.append('ELSE')
        L.append('BEGIN')
        L.append('    UPDATE #IMP_OPT SET PREREQ_OK = 1;')
        L.append("    PRINT 'Prerequisites OK on ' + DB_NAME() + CASE WHEN @OLD IS NOT NULL THEN ' (month ' + @YM + ' ID ' + CONVERT(VARCHAR(10), @OLD) + ' will be replaced)' ELSE '' END + '.';")
        L.append('END')
        L.append('GO')
        L.append('')
        # 2. staging
        L.append('/* ---- 2. staging (#temp tables only) ------------------------------------------------------- */')
        L.append("IF OBJECT_ID('tempdb..#IMP_M') IS NOT NULL DROP TABLE #IMP_M;")
        L.append("IF OBJECT_ID('tempdb..#IMP_W') IS NOT NULL DROP TABLE #IMP_W;")
        L.append("IF OBJECT_ID('tempdb..#IMP_G') IS NOT NULL DROP TABLE #IMP_G;")
        L.append("IF OBJECT_ID('tempdb..#IMP_T') IS NOT NULL DROP TABLE #IMP_T;")
        L.append('CREATE TABLE #IMP_M (TITLE NVARCHAR(100) NULL, TABS NVARCHAR(MAX) NOT NULL);')
        cols = ',\n'.join('    %-20s NVARCHAR(%d) NULL' % (k, reg.maxlen[k]) for k in reg.keys)
        L.append('CREATE TABLE #IMP_W (SEQ INT NOT NULL PRIMARY KEY, TAB_CODE VARCHAR(20) NOT NULL, SHEET_ROW INT NOT NULL, SORT_ORDER FLOAT NOT NULL,')
        L.append('    STATUS_CODE VARCHAR(20) NULL, FF_KEY VARCHAR(25) NULL, MATCH_EXPTNO VARCHAR(25) NULL, LOG_TEXT NVARCHAR(400) NOT NULL,')
        L.append(cols + ',')
        L.append('    CELL_FMT NVARCHAR(MAX) NULL, UNIQUE (TAB_CODE, SORT_ORDER));')
        L.append('CREATE TABLE #IMP_G (SEQ INT NOT NULL PRIMARY KEY, TAB_CODE VARCHAR(20) NOT NULL, SHEET_ROW INT NOT NULL, SORT_ORDER FLOAT NOT NULL,')
        L.append('    DATA NVARCHAR(MAX) NULL, CELL_FMT NVARCHAR(MAX) NULL, LOG_TEXT NVARCHAR(400) NOT NULL, UNIQUE (TAB_CODE, SORT_ORDER));')
        L.append('CREATE TABLE #IMP_T (TAB_CODE VARCHAR(20) NOT NULL PRIMARY KEY, META NVARCHAR(MAX) NOT NULL);')
        L.append('INSERT INTO #IMP_M (TITLE, TABS) VALUES (%s,\n    %s);' % (sql_n(a.title), sql_n(imp.tabs_json)))
        for code, mj in imp.meta:
            L.append('INSERT INTO #IMP_T (TAB_CODE, META) VALUES (%s, %s);' % (sql_a(code), sql_n(mj)))
        L.append('GO')
        cur = None
        for x in imp.wo:
            if x['tab'] != cur:
                if cur is not None:
                    L.append('GO')
                cur = x['tab']
                L.append('-- %s' % cur)
            keys = [k for k in reg.keys if k in x['values']]
            names = ['SEQ', 'TAB_CODE', 'SHEET_ROW', 'SORT_ORDER', 'STATUS_CODE', 'FF_KEY', 'LOG_TEXT'] + keys + ['CELL_FMT']
            vals = [str(x['seq']), sql_a(x['tab']), str(x['sheet_row']), '%.1f' % x['sort'], sql_a(x['status']), sql_a(x['ff_key']),
                    sql_n(x['log'])] + [sql_n(x['values'][k]) for k in keys] + [sql_n(x['fmt'])]
            L.append('INSERT INTO #IMP_W (%s)\nVALUES (%s);' % (', '.join(names), ', '.join(vals)))
        if imp.wo:
            L.append('GO')
        cur = None
        for i, g in enumerate(imp.grid, 1):
            if g['tab'] != cur:
                if cur is not None:
                    L.append('GO')
                cur = g['tab']
                L.append('-- %s' % cur)
            L.append('INSERT INTO #IMP_G (SEQ, TAB_CODE, SHEET_ROW, SORT_ORDER, DATA, CELL_FMT, LOG_TEXT) VALUES (%d, %s, %d, %.1f, %s, %s, %s);'
                     % (i, sql_a(g['tab']), g['sheet_row'], g['sort'], sql_n(g['data']), sql_n(g['fmt']), sql_n(g['log'])))
        if imp.grid:
            L.append('GO')
        L.append('')
        # 3. apply
        exp_w = [(t['code'], n) for t, _, n in imp.wo_tabs]
        exp_g = sorted(grid_tab.items())
        L.append('/* ---- 3. import: one transaction ----------------------------------------------------------- */')
        L.append('SET XACT_ABORT ON;')
        L.append('DECLARE @CMP VARCHAR(6), @CITY VARCHAR(3), @YM CHAR(7), @U VARCHAR(50), @UN NVARCHAR(100), @NOW DATETIME, @MID INT, @OLD INT,')
        L.append('        @REPLACE BIT, @MIRROR BIT, @EXPTNO VARCHAR(25), @IDS VARCHAR(MAX), @MSG NVARCHAR(2048), @N INT, @LINKED INT,')
        L.append('        @NW INT, @NL INT, @NG INT, @NT INT;')
        L.append('DECLARE @problems TABLE (n INT IDENTITY(1,1), problem NVARCHAR(1000));')
        L.append('DECLARE @MAPW TABLE (ID INT NOT NULL PRIMARY KEY, TAB_CODE VARCHAR(20) NOT NULL, SORT_ORDER FLOAT NOT NULL, EXPTNO VARCHAR(25) NULL);')
        L.append('DECLARE @MAPG TABLE (ID INT NOT NULL PRIMARY KEY, TAB_CODE VARCHAR(20) NOT NULL, SORT_ORDER FLOAT NOT NULL);')
        L.append('SELECT @CMP = %s, @CITY = %s, @YM = %s, @U = %s, @UN = %s, @LINKED = 0;'
                 % (sql_a(self.cmp), sql_a(self.city), sql_a(self.ym), sql_a(IMPORT_USER), sql_n(IMPORT_USER)))
        L.append("IF UPPER(DB_NAME()) <> 'LINKSDB20'")
        L.append("BEGIN RAISERROR('Not on LinksDB20 - nothing imported.', 16, 1); RETURN; END")
        L.append("IF OBJECT_ID('tempdb..#IMP_OPT') IS NULL OR NOT EXISTS (SELECT 1 FROM #IMP_OPT WHERE PREREQ_OK = 1)")
        L.append("BEGIN RAISERROR('Step 1 did not pass (or did not run) - nothing imported.', 16, 1); RETURN; END")
        L.append('SELECT @REPLACE = REPLACE_MONTH, @MIRROR = MIRROR_TRACKING FROM #IMP_OPT;')
        L.append('-- every staging row arrived (a failed INSERT in step 2 shows up here, before anything is written)')
        L.append('INSERT INTO @problems (problem)')
        L.append("SELECT N'#IMP_W ' + x.t + N': ' + CONVERT(NVARCHAR(10), ISNULL(s.n, 0)) + N' row(s) staged, ' + CONVERT(NVARCHAR(10), x.n) + N' expected.'")
        L.append('  FROM (VALUES %s) x (t, n)' % (', '.join("('%s', %d)" % e for e in exp_w) or "('-', 0)"))
        L.append('  LEFT JOIN (SELECT TAB_CODE, COUNT(*) AS n FROM #IMP_W GROUP BY TAB_CODE) s ON s.TAB_CODE = x.t')
        L.append(' WHERE ISNULL(s.n, 0) <> x.n;')
        L.append('INSERT INTO @problems (problem)')
        L.append("SELECT N'#IMP_G ' + x.t + N': ' + CONVERT(NVARCHAR(10), ISNULL(s.n, 0)) + N' row(s) staged, ' + CONVERT(NVARCHAR(10), x.n) + N' expected.'")
        L.append('  FROM (VALUES %s) x (t, n)' % (', '.join("('%s', %d)" % e for e in exp_g) or "('-', 0)"))
        L.append('  LEFT JOIN (SELECT TAB_CODE, COUNT(*) AS n FROM #IMP_G GROUP BY TAB_CODE) s ON s.TAB_CODE = x.t')
        L.append(' WHERE ISNULL(s.n, 0) <> x.n;')
        L.append("IF (SELECT COUNT(*) FROM #IMP_W) <> %d INSERT INTO @problems (problem) VALUES (N'#IMP_W holds rows of another tab.');" % n_wo)
        L.append("IF (SELECT COUNT(*) FROM #IMP_G) <> %d INSERT INTO @problems (problem) VALUES (N'#IMP_G holds rows of another tab.');" % n_g)
        L.append("IF (SELECT COUNT(*) FROM #IMP_T) <> %d INSERT INTO @problems (problem) VALUES (N'#IMP_T: tab layouts missing.');" % n_t)
        L.append("IF (SELECT COUNT(*) FROM #IMP_M) <> 1 INSERT INTO @problems (problem) VALUES (N'#IMP_M: month row missing.');")
        L.append('IF EXISTS (SELECT 1 FROM @problems)')
        L.append('BEGIN')
        L.append('    SELECT n AS [#], problem FROM @problems ORDER BY n;')
        L.append("    RAISERROR('Staging is incomplete - nothing imported.', 16, 1);")
        L.append('    RETURN;')
        L.append('END')
        L.append('')
        L.append('BEGIN TRY')
        L.append('    BEGIN TRANSACTION;')
        L.append('    SET @NOW = GETDATE();')
        L.append('    -- the month again, under lock (two runs at once, or someone created it since step 1)')
        L.append('    SELECT @OLD = MIN(ID) FROM dbo.TBL_EXPT_SHEET_MONTH WITH (UPDLOCK, HOLDLOCK)')
        L.append('     WHERE CMPCODE = @CMP AND CITYCODE = @CITY AND SHEET_MONTH = @YM;')
        L.append('    IF @OLD IS NOT NULL AND @REPLACE = 0')
        L.append("        RAISERROR('The month exists (created after step 1?) and @Replace = 0.', 16, 1);")
        L.append('    IF @OLD IS NOT NULL')
        L.append('    BEGIN')
        L.append('        SELECT @N = COUNT(*) FROM dbo.TBL_EXPT_SHEET_ROW WITH (UPDLOCK, HOLDLOCK) WHERE MONTH_ID = @OLD;')
        L.append('        SELECT @N = COUNT(*) FROM dbo.TBL_EXPT_SHEET_GRIDROW WITH (UPDLOCK, HOLDLOCK) WHERE MONTH_ID = @OLD;')
        self.pristine_checks(L, '@OLD', 'Cannot replace month ' + self.ym)
        L.append('        IF EXISTS (SELECT 1 FROM @problems)')
        L.append('        BEGIN')
        L.append("            SELECT @MSG = STUFF((SELECT N' | ' + problem FROM @problems ORDER BY n FOR XML PATH(''), TYPE).value('.', 'NVARCHAR(MAX)'), 1, 3, N'');")
        L.append("            RAISERROR(N'%s', 16, 1, @MSG);")
        L.append('        END')
        L.append("        DELETE L FROM dbo.TBL_EXPT_SHEET_LOG L")
        L.append("         WHERE (L.ROW_KIND = 'M' AND L.ROW_ID = @OLD)")
        L.append("            OR (L.ROW_KIND = 'W' AND L.ROW_ID IN (SELECT ID FROM dbo.TBL_EXPT_SHEET_ROW WHERE MONTH_ID = @OLD))")
        L.append("            OR (L.ROW_KIND = 'G' AND L.ROW_ID IN (SELECT ID FROM dbo.TBL_EXPT_SHEET_GRIDROW WHERE MONTH_ID = @OLD));")
        L.append('        DELETE FROM dbo.TBL_EXPT_SHEET_ROW WHERE MONTH_ID = @OLD;')
        L.append('        DELETE FROM dbo.TBL_EXPT_SHEET_GRIDROW WHERE MONTH_ID = @OLD;')
        L.append('        DELETE FROM dbo.TBL_EXPT_SHEET_TABMETA WHERE MONTH_ID = @OLD;')
        L.append('        DELETE FROM dbo.TBL_EXPT_SHEET_MONTH WHERE ID = @OLD;')
        L.append("        PRINT 'Earlier import of the month (ID ' + CONVERT(VARCHAR(10), @OLD) + ') removed.';")
        L.append('    END')
        L.append('')
        L.append('    -- the month')
        L.append('    INSERT INTO dbo.TBL_EXPT_SHEET_MONTH (CMPCODE, CITYCODE, SHEET_MONTH, TITLE, TABS, ROWVER, MAKER_ID, MAKER_DT, MODIFY_ID, MODIFY_DT)')
        L.append('    SELECT @CMP, @CITY, @YM, TITLE, TABS, 1, @U, @NOW, @U, @NOW FROM #IMP_M;')
        L.append('    SET @MID = SCOPE_IDENTITY();')
        L.append('    INSERT INTO dbo.TBL_EXPT_SHEET_LOG (ROW_KIND, ROW_ID, COL_KEY, OLD_VALUE, NEW_VALUE, ACTION, USER_ID, USER_NAME, LOG_DT)')
        L.append("    VALUES ('M', @MID, NULL, NULL, LEFT(@YM + N' imported from workbook ' + %s, 2000), 'insert', @U, @UN, @NOW);"
                 % sql_n('"%s" (%s)' % (os.path.basename(a.workbook), self.wb_hash[:12])))
        L.append('')
        L.append('    -- FF JOB -> EXPTNO: a job of the same company / branch that is not cancelled')
        L.append('    UPDATE I SET MATCH_EXPTNO = I.FF_KEY')
        L.append('      FROM #IMP_W I')
        L.append('     WHERE I.FF_KEY IS NOT NULL')
        L.append("       AND EXISTS (SELECT 1 FROM dbo.EXPT_MASTER E WHERE E.EXPTNO = I.FF_KEY AND LTRIM(RTRIM(ISNULL(E.CMPCODE, ''))) = @CMP")
        L.append("                   AND LTRIM(RTRIM(ISNULL(E.CITYCODE, ''))) = @CITY AND ISNULL(E.EXPTCANCEL, '') <> 'Y');")
        L.append('')
        L.append('    -- WO rows, in workbook order (identity follows the ORDER BY)')
        cl = ', '.join(reg.keys)
        L.append('    INSERT INTO dbo.TBL_EXPT_SHEET_ROW (MONTH_ID, TAB_CODE, SORT_ORDER, STATUS_CODE, EXPTNO,')
        L.append('           ' + cl + ',')
        L.append('           CELL_FMT, ROWVER, ISDELETED, MAKER_ID, MAKER_DT, MODIFY_ID, MODIFY_DT)')
        L.append('    OUTPUT INSERTED.ID, INSERTED.TAB_CODE, INSERTED.SORT_ORDER, INSERTED.EXPTNO INTO @MAPW (ID, TAB_CODE, SORT_ORDER, EXPTNO)')
        L.append('    SELECT @MID, I.TAB_CODE, I.SORT_ORDER, I.STATUS_CODE, CASE WHEN @MIRROR = 0 THEN I.MATCH_EXPTNO END,')
        L.append('           ' + ', '.join('I.' + k for k in reg.keys) + ',')
        L.append('           I.CELL_FMT, 1, 0, @U, @NOW, @U, @NOW')
        L.append('      FROM #IMP_W I')
        L.append('     ORDER BY I.SEQ;')
        L.append('    INSERT INTO dbo.TBL_EXPT_SHEET_LOG (ROW_KIND, ROW_ID, COL_KEY, OLD_VALUE, NEW_VALUE, ACTION, USER_ID, USER_NAME, LOG_DT)')
        L.append("    SELECT 'W', M.ID, NULL, NULL, I.LOG_TEXT + CASE WHEN I.MATCH_EXPTNO IS NOT NULL THEN N'; FF JOB is job ' + I.MATCH_EXPTNO ELSE N'' END,")
        L.append("           'insert', @U, @UN, @NOW")
        L.append('      FROM @MAPW M JOIN #IMP_W I ON I.TAB_CODE = M.TAB_CODE AND I.SORT_ORDER = M.SORT_ORDER')
        L.append('     ORDER BY M.ID;')
        L.append('')
        L.append('    -- grid / free-form rows and tab layouts')
        L.append('    INSERT INTO dbo.TBL_EXPT_SHEET_GRIDROW (MONTH_ID, TAB_CODE, SORT_ORDER, DATA, CELL_FMT, ROWVER, ISDELETED, MAKER_ID, MAKER_DT, MODIFY_ID, MODIFY_DT)')
        L.append('    OUTPUT INSERTED.ID, INSERTED.TAB_CODE, INSERTED.SORT_ORDER INTO @MAPG (ID, TAB_CODE, SORT_ORDER)')
        L.append('    SELECT @MID, G.TAB_CODE, G.SORT_ORDER, G.DATA, G.CELL_FMT, 1, 0, @U, @NOW, @U, @NOW FROM #IMP_G G ORDER BY G.SEQ;')
        L.append('    INSERT INTO dbo.TBL_EXPT_SHEET_LOG (ROW_KIND, ROW_ID, COL_KEY, OLD_VALUE, NEW_VALUE, ACTION, USER_ID, USER_NAME, LOG_DT)')
        L.append("    SELECT 'G', M.ID, NULL, NULL, G.LOG_TEXT, 'insert', @U, @UN, @NOW")
        L.append('      FROM @MAPG M JOIN #IMP_G G ON G.TAB_CODE = M.TAB_CODE AND G.SORT_ORDER = M.SORT_ORDER')
        L.append('     ORDER BY M.ID;')
        L.append('    INSERT INTO dbo.TBL_EXPT_SHEET_TABMETA (MONTH_ID, TAB_CODE, META, ROWVER, MODIFY_ID, MODIFY_DT)')
        L.append('    SELECT @MID, T.TAB_CODE, T.META, 1, @U, @NOW FROM #IMP_T T;')
        L.append('    INSERT INTO dbo.TBL_EXPT_SHEET_LOG (ROW_KIND, ROW_ID, COL_KEY, OLD_VALUE, NEW_VALUE, ACTION, USER_ID, USER_NAME, LOG_DT)')
        L.append("    SELECT 'M', @MID, LEFT('META.' + T.TAB_CODE, 40), NULL, LEFT(T.META, 2000), 'insert', @U, @UN, @NOW FROM #IMP_T T ORDER BY T.TAB_CODE;")
        L.append('')
        L.append('    -- @MirrorTracking = 1: every matched job is linked by dbo.USP_EXPT_SHEET_LINK (EXPTNO + FF_JOB, the tracking mirror,')
        L.append("    -- its own 'link' history entries); each call returns its own result sets")
        L.append('    IF @MIRROR = 1')
        L.append('    BEGIN')
        L.append('        DECLARE imp_link CURSOR LOCAL FAST_FORWARD FOR')
        L.append('            SELECT DISTINCT I.MATCH_EXPTNO FROM #IMP_W I WHERE I.MATCH_EXPTNO IS NOT NULL ORDER BY I.MATCH_EXPTNO;')
        L.append('        OPEN imp_link;')
        L.append('        FETCH NEXT FROM imp_link INTO @EXPTNO;')
        L.append('        WHILE @@FETCH_STATUS = 0')
        L.append('        BEGIN')
        L.append("            SET @IDS = STUFF((SELECT ',' + CONVERT(VARCHAR(12), M.ID) FROM @MAPW M JOIN #IMP_W I ON I.TAB_CODE = M.TAB_CODE AND I.SORT_ORDER = M.SORT_ORDER")
        L.append("                              WHERE I.MATCH_EXPTNO = @EXPTNO ORDER BY M.ID FOR XML PATH('')), 1, 1, '');")
        L.append('            EXEC dbo.USP_EXPT_SHEET_LINK @IDS, @EXPTNO, @U, @UN;')
        L.append('            IF @@TRANCOUNT = 0')
        L.append("                RAISERROR('dbo.USP_EXPT_SHEET_LINK rolled the transaction back.', 16, 1);")
        L.append('            IF EXISTS (SELECT 1 FROM dbo.TBL_EXPT_SHEET_ROW R JOIN @MAPW M ON M.ID = R.ID')
        L.append('                        JOIN #IMP_W I ON I.TAB_CODE = M.TAB_CODE AND I.SORT_ORDER = M.SORT_ORDER')
        L.append("                       WHERE I.MATCH_EXPTNO = @EXPTNO AND ISNULL(R.EXPTNO, '') <> @EXPTNO)")
        L.append("                RAISERROR('dbo.USP_EXPT_SHEET_LINK did not link the rows of job %s (see its STATUS row above).', 16, 1, @EXPTNO);")
        L.append('            SET @LINKED = @LINKED + 1;')
        L.append('            FETCH NEXT FROM imp_link INTO @EXPTNO;')
        L.append('        END')
        L.append('        CLOSE imp_link;')
        L.append('        DEALLOCATE imp_link;')
        L.append('    END')
        L.append('')
        L.append('    -- last look before the commit')
        L.append('    IF (SELECT COUNT(*) FROM dbo.TBL_EXPT_SHEET_ROW WHERE MONTH_ID = @MID) <> %d' % n_wo)
        L.append("       OR (SELECT COUNT(*) FROM dbo.TBL_EXPT_SHEET_GRIDROW WHERE MONTH_ID = @MID) <> %d" % n_g)
        L.append("       OR (SELECT COUNT(*) FROM dbo.TBL_EXPT_SHEET_TABMETA WHERE MONTH_ID = @MID) <> %d" % n_t)
        L.append("       OR (SELECT COUNT(*) FROM dbo.TBL_EXPT_SHEET_ROW WHERE MONTH_ID = @MID AND EXPTNO IS NOT NULL)")
        L.append('          <> (SELECT COUNT(*) FROM #IMP_W WHERE MATCH_EXPTNO IS NOT NULL)')
        L.append("        RAISERROR('Row counts after the insert do not match the staging - rolled back.', 16, 1);")
        L.append('    COMMIT TRANSACTION;')
        L.append('    SELECT @NW = (SELECT COUNT(*) FROM @MAPW), @NL = (SELECT COUNT(*) FROM #IMP_W WHERE MATCH_EXPTNO IS NOT NULL),')
        L.append('           @NG = (SELECT COUNT(*) FROM @MAPG), @NT = (SELECT COUNT(*) FROM #IMP_T);')
        L.append("    PRINT 'Imported month ' + @YM + ' as ID ' + CONVERT(VARCHAR(10), @MID) + ': ' + CONVERT(VARCHAR(10), @NW) + ' WO rows ('")
        L.append("          + CONVERT(VARCHAR(10), @NL) + ' linked'")
        L.append("          + CASE WHEN @MIRROR = 1 THEN ' through USP_EXPT_SHEET_LINK, ' + CONVERT(VARCHAR(10), @LINKED) + ' job(s)' ELSE '' END + '), '")
        L.append("          + CONVERT(VARCHAR(10), @NG) + ' grid rows, ' + CONVERT(VARCHAR(10), @NT) + ' tab layouts.';")
        L.append('END TRY')
        L.append('BEGIN CATCH')
        L.append('    IF @@TRANCOUNT > 0 ROLLBACK TRANSACTION;')
        L.append('    SET @MSG = ERROR_MESSAGE();')
        L.append("    RAISERROR(N'*** Import failed and was rolled back - nothing was written: %s', 16, 1, @MSG);")
        L.append('END CATCH')
        L.append('GO')
        L.append('SET NOEXEC OFF;')
        L.append('GO')
        L.append('')
        # 4. verify
        L.append('/* ---- 4. verify (read-only) ----------------------------------------------------------------- */')
        L.append("IF OBJECT_ID('dbo.TBL_EXPT_SHEET_MONTH', 'U') IS NULL OR OBJECT_ID('dbo.TBL_EXPT_SHEET_ROW', 'U') IS NULL")
        L.append("    PRINT '*** The sheet tables are not there - nothing to verify.';")
        L.append('ELSE')
        L.append('BEGIN')
        L.append('    DECLARE @MID INT;')
        L.append('    SELECT @MID = MIN(ID) FROM dbo.TBL_EXPT_SHEET_MONTH WHERE CMPCODE = %s AND CITYCODE = %s AND SHEET_MONTH = %s;'
                 % (sql_a(self.cmp), sql_a(self.city), sql_a(self.ym)))
        L.append('    SELECT ID, CMPCODE, CITYCODE, SHEET_MONTH, TITLE, ROWVER, LEN(TABS) AS TABS_CHARS, MAKER_ID, MAKER_DT FROM dbo.TBL_EXPT_SHEET_MONTH WHERE ID = @MID;')
        L.append('    -- WO rows per tab: expected (from the workbook) vs in the table')
        L.append("    SELECT x.TAB_CODE, x.EXPECTED, ISNULL(r.ROWS_IN, 0) AS IMPORTED, ISNULL(r.LINKED, 0) AS LINKED, ISNULL(r.FF_TEXT_ONLY, 0) AS FF_JOB_NOT_LINKED,")
        L.append("           ISNULL(r.WITH_STATUS, 0) AS WITH_STATUS, ISNULL(r.WITH_FMT, 0) AS WITH_CELL_FMT,")
        L.append("           CASE WHEN ISNULL(r.ROWS_IN, 0) = x.EXPECTED THEN 'ok' ELSE '*** MISMATCH' END AS CHECK_RESULT")
        L.append('      FROM (VALUES %s) x (TAB_CODE, EXPECTED)' % (', '.join("('%s', %d)" % e for e in exp_w) or "('-', 0)"))
        L.append('      LEFT JOIN (SELECT TAB_CODE, COUNT(*) AS ROWS_IN, SUM(CASE WHEN EXPTNO IS NOT NULL THEN 1 ELSE 0 END) AS LINKED,')
        L.append("                        SUM(CASE WHEN EXPTNO IS NULL AND LTRIM(RTRIM(ISNULL(FF_JOB, N''))) <> N'' THEN 1 ELSE 0 END) AS FF_TEXT_ONLY,")
        L.append('                        SUM(CASE WHEN STATUS_CODE IS NOT NULL THEN 1 ELSE 0 END) AS WITH_STATUS, SUM(CASE WHEN CELL_FMT IS NOT NULL THEN 1 ELSE 0 END) AS WITH_FMT')
        L.append('                   FROM dbo.TBL_EXPT_SHEET_ROW WHERE MONTH_ID = @MID AND ISDELETED = 0 GROUP BY TAB_CODE) r ON r.TAB_CODE = x.TAB_CODE')
        L.append('     ORDER BY x.TAB_CODE;')
        L.append("    SELECT TAB_CODE, ISNULL(STATUS_CODE, '(none)') AS STATUS_CODE, COUNT(*) AS ROWS_IN FROM dbo.TBL_EXPT_SHEET_ROW")
        L.append('     WHERE MONTH_ID = @MID AND ISDELETED = 0 GROUP BY TAB_CODE, STATUS_CODE ORDER BY TAB_CODE, STATUS_CODE;')
        L.append('    -- grid rows and tab layouts')
        L.append("    SELECT x.TAB_CODE, x.EXPECTED, ISNULL(g.ROWS_IN, 0) AS IMPORTED, CASE WHEN t.TAB_CODE IS NULL THEN 'no' ELSE 'yes' END AS TAB_LAYOUT,")
        L.append("           CASE WHEN ISNULL(g.ROWS_IN, 0) = x.EXPECTED AND t.TAB_CODE IS NOT NULL THEN 'ok' ELSE '*** MISMATCH' END AS CHECK_RESULT")
        L.append('      FROM (VALUES %s) x (TAB_CODE, EXPECTED)' % (', '.join("('%s', %d)" % (c, grid_tab.get(c, 0)) for c, _ in imp.meta) or "('-', 0)"))
        L.append('      LEFT JOIN (SELECT TAB_CODE, COUNT(*) AS ROWS_IN FROM dbo.TBL_EXPT_SHEET_GRIDROW WHERE MONTH_ID = @MID AND ISDELETED = 0 GROUP BY TAB_CODE) g ON g.TAB_CODE = x.TAB_CODE')
        L.append('      LEFT JOIN dbo.TBL_EXPT_SHEET_TABMETA t ON t.MONTH_ID = @MID AND t.TAB_CODE = x.TAB_CODE')
        L.append('     ORDER BY x.TAB_CODE;')
        L.append('    -- history entries written by the import')
        L.append("    SELECT L.ROW_KIND, L.ACTION, COUNT(*) AS ENTRIES FROM dbo.TBL_EXPT_SHEET_LOG L")
        L.append("     WHERE L.USER_ID = %s AND ((L.ROW_KIND = 'M' AND L.ROW_ID = @MID)" % sql_a(IMPORT_USER))
        L.append("        OR (L.ROW_KIND = 'W' AND L.ROW_ID IN (SELECT ID FROM dbo.TBL_EXPT_SHEET_ROW WHERE MONTH_ID = @MID))")
        L.append("        OR (L.ROW_KIND = 'G' AND L.ROW_ID IN (SELECT ID FROM dbo.TBL_EXPT_SHEET_GRIDROW WHERE MONTH_ID = @MID)))")
        L.append('     GROUP BY L.ROW_KIND, L.ACTION ORDER BY L.ROW_KIND, L.ACTION;')
        L.append("    SELECT 'expected' AS NOTE, %d AS MONTH_INSERT, %d AS WO_INSERT, %d AS GRID_INSERT, %d AS TAB_LAYOUT_INSERT;" % (1 + n_t, n_wo, n_g, n_t))
        L.append('    -- FF JOB values that stayed text (not linked), with the reason')
        L.append('    SELECT R.TAB_CODE, R.ID, R.FF_JOB, ' + self.match_case("CONVERT(VARCHAR(25), CASE WHEN LEN(LTRIM(RTRIM(R.FF_JOB))) <= 25 THEN LTRIM(RTRIM(R.FF_JOB)) END)",
                                                                   sql_a(self.cmp), sql_a(self.city)) + ' AS REASON')
        L.append('      FROM dbo.TBL_EXPT_SHEET_ROW R')
        L.append("     WHERE R.MONTH_ID = @MID AND R.ISDELETED = 0 AND R.EXPTNO IS NULL AND LTRIM(RTRIM(ISNULL(R.FF_JOB, N''))) <> N''")
        L.append('     ORDER BY R.TAB_CODE, R.SORT_ORDER;')
        L.append('END')
        L.append("IF OBJECT_ID('tempdb..#IMP_W') IS NOT NULL DROP TABLE #IMP_W;")
        L.append("IF OBJECT_ID('tempdb..#IMP_G') IS NOT NULL DROP TABLE #IMP_G;")
        L.append("IF OBJECT_ID('tempdb..#IMP_T') IS NOT NULL DROP TABLE #IMP_T;")
        L.append("IF OBJECT_ID('tempdb..#IMP_M') IS NOT NULL DROP TABLE #IMP_M;")
        L.append("IF OBJECT_ID('tempdb..#IMP_OPT') IS NOT NULL DROP TABLE #IMP_OPT;")
        L.append('GO')
        return L

    # -- the rollback -----------------------------------------------------------------------------------------------
    def rollback_script(self):
        a = self.a
        L = []
        u = sql_a(IMPORT_USER)
        self.header(L, self.names['rollback'], [
            'Rollback for %s: removes exactly that import of month %s (CMPCODE %s / CITYCODE %s,' % (self.names['import'], self.ym, self.cmp, self.city),
            '"%s") - its WO rows, grid rows, tab layouts, history entries and the month itself.' % ascii_comment(a.title),
        ], [
            'Refuses (and changes nothing) when the month was not created by the import (MAKER_ID \'%s\'), or when anyone' % IMPORT_USER,
            'has touched it since: a WO row, grid row, tab layout or the month itself with MAKER_ID / MODIFY_ID other than the',
            'import, a history entry by another user (edits, formats, moves, links, deletes, restores), or a personal view on',
            'the month. Then the Sheet view is in use: fix it there instead.',
            '',
            'NOT undone: EXPT_JOB_TRACKING / EXPT_JOB_TRACKING_LOG rows written when the import ran with @MirrorTracking = 1',
            '(step 1 lists them) - they are the tracking module\'s ordinary data.',
            'After it: EXPORT_SHEET_VIEW\'s own rollback can drop the (then empty) sheet tables if wanted.',
            '',
            'Safety : step 1 only reads; step 2 deletes in ONE transaction after the same checks under lock; step 3 only reads.',
            '         Stops on any database other than LinksDB20. Re-runnable (the second run finds nothing to remove).',
            'Run as : a login with DELETE on the sheet tables. SQL Server 2008 R2.',
        ])
        L.append('SET NOEXEC OFF;')
        L.append('GO')
        L.append('USE [LinksDB20]')
        L.append('GO')
        L.append('SET NOCOUNT ON;')
        L.append('SET XACT_ABORT ON;')
        L.append('GO')
        L.append('')
        L.append('/* ---- 1. prerequisites (read-only) ---------------------------------------------------------- */')
        L.append('DECLARE @problems TABLE (n INT IDENTITY(1,1), problem NVARCHAR(1000));')
        L.append("IF UPPER(DB_NAME()) <> 'LINKSDB20'")
        L.append("    INSERT INTO @problems (problem) VALUES (N'Connected to ' + DB_NAME() + N', not LinksDB20.');")
        L.append('INSERT INTO @problems (problem)')
        L.append("SELECT N'dbo.' + x.t + N' not found - nothing to roll back.'")
        L.append("  FROM (VALUES ('TBL_EXPT_SHEET_MONTH'), ('TBL_EXPT_SHEET_ROW'), ('TBL_EXPT_SHEET_GRIDROW'), ('TBL_EXPT_SHEET_TABMETA'),")
        L.append("               ('TBL_EXPT_SHEET_LOG'), ('TBL_EXPT_SHEET_VIEW')) x (t)")
        L.append(" WHERE OBJECT_ID('dbo.' + x.t, 'U') IS NULL;")
        L.append('INSERT INTO @problems (problem)')
        L.append("SELECT N'dbo.' + x.t + N'.' + x.c + N' not found.'")
        L.append("  FROM (VALUES ('TBL_EXPT_SHEET_MONTH', 'MAKER_ID'), ('TBL_EXPT_SHEET_MONTH', 'MODIFY_ID'), ('TBL_EXPT_SHEET_ROW', 'MAKER_ID'),")
        L.append("               ('TBL_EXPT_SHEET_ROW', 'MODIFY_ID'), ('TBL_EXPT_SHEET_GRIDROW', 'MAKER_ID'), ('TBL_EXPT_SHEET_GRIDROW', 'MODIFY_ID'),")
        L.append("               ('TBL_EXPT_SHEET_TABMETA', 'MODIFY_ID'), ('TBL_EXPT_SHEET_LOG', 'USER_ID'), ('TBL_EXPT_SHEET_VIEW', 'MONTH_ID')) x (t, c)")
        L.append(" WHERE OBJECT_ID('dbo.' + x.t, 'U') IS NOT NULL AND COL_LENGTH('dbo.' + x.t, x.c) IS NULL;")
        L.append('INSERT INTO @problems (problem)')
        L.append("SELECT N'No DELETE permission on dbo.' + x.t + N'.'")
        L.append("  FROM (VALUES ('TBL_EXPT_SHEET_MONTH'), ('TBL_EXPT_SHEET_ROW'), ('TBL_EXPT_SHEET_GRIDROW'), ('TBL_EXPT_SHEET_TABMETA'), ('TBL_EXPT_SHEET_LOG')) x (t)")
        L.append(" WHERE OBJECT_ID('dbo.' + x.t, 'U') IS NOT NULL AND ISNULL(HAS_PERMS_BY_NAME('dbo.' + x.t, 'OBJECT', 'DELETE'), 0) = 0;")
        L.append('IF EXISTS (SELECT 1 FROM @problems)')
        L.append('BEGIN')
        L.append("    PRINT '*** Rollback check failed - NOTHING was removed:';")
        L.append('    SELECT n AS [#], problem FROM @problems ORDER BY n;')
        L.append("    RAISERROR('Rollback check failed - see the list above.', 16, 1);")
        L.append('    SET NOEXEC ON;')
        L.append('END')
        L.append('GO')
        L.append('')
        L.append('DECLARE @problems TABLE (n INT IDENTITY(1,1), problem NVARCHAR(1000));')
        L.append('DECLARE @MID INT;')
        L.append('SELECT @MID = MIN(ID) FROM dbo.TBL_EXPT_SHEET_MONTH WHERE CMPCODE = %s AND CITYCODE = %s AND SHEET_MONTH = %s;'
                 % (sql_a(self.cmp), sql_a(self.city), sql_a(self.ym)))
        L.append('IF @MID IS NULL')
        L.append('BEGIN')
        L.append("    PRINT 'Month %s is not there for %s / %s - nothing to roll back.';" % (self.ym, self.cmp, self.city))
        L.append('    SET NOEXEC ON;')
        L.append('END')
        L.append('ELSE')
        L.append('BEGIN')
        L.append("    IF NOT EXISTS (SELECT 1 FROM dbo.TBL_EXPT_SHEET_MONTH WHERE ID = @MID AND MAKER_ID = %s)" % u)
        L.append("        INSERT INTO @problems (problem) VALUES (N'Month ID ' + CONVERT(NVARCHAR(10), @MID) + N' was not created by the sheet import - it is left alone.');")
        self.pristine_checks(L, '@MID', 'Month ' + self.ym)
        L.append('    -- what will go')
        L.append("    SELECT 'month' AS WHAT, COUNT(*) AS ROWS_TO_REMOVE FROM dbo.TBL_EXPT_SHEET_MONTH WHERE ID = @MID")
        L.append("    UNION ALL SELECT 'WO rows', COUNT(*) FROM dbo.TBL_EXPT_SHEET_ROW WHERE MONTH_ID = @MID")
        L.append("    UNION ALL SELECT 'grid rows', COUNT(*) FROM dbo.TBL_EXPT_SHEET_GRIDROW WHERE MONTH_ID = @MID")
        L.append("    UNION ALL SELECT 'tab layouts', COUNT(*) FROM dbo.TBL_EXPT_SHEET_TABMETA WHERE MONTH_ID = @MID")
        L.append("    UNION ALL SELECT 'history entries', COUNT(*) FROM dbo.TBL_EXPT_SHEET_LOG L WHERE (L.ROW_KIND = 'M' AND L.ROW_ID = @MID)")
        L.append("              OR (L.ROW_KIND = 'W' AND L.ROW_ID IN (SELECT ID FROM dbo.TBL_EXPT_SHEET_ROW WHERE MONTH_ID = @MID))")
        L.append("              OR (L.ROW_KIND = 'G' AND L.ROW_ID IN (SELECT ID FROM dbo.TBL_EXPT_SHEET_GRIDROW WHERE MONTH_ID = @MID));")
        L.append('    -- tracking written by @MirrorTracking = 1 (kept)')
        L.append("    IF OBJECT_ID('dbo.EXPT_JOB_TRACKING_LOG', 'U') IS NOT NULL AND COL_LENGTH('dbo.EXPT_JOB_TRACKING_LOG', 'CHANGED_BY') IS NOT NULL")
        L.append("        SELECT 'kept: EXPT_JOB_TRACKING_LOG entries by the import' AS NOTE, COUNT(*) AS ENTRIES, COUNT(DISTINCT T.EXPTNO) AS JOBS")
        L.append('          FROM dbo.EXPT_JOB_TRACKING_LOG T')
        L.append('         WHERE T.CHANGED_BY = %s AND T.EXPTNO IN (SELECT EXPTNO FROM dbo.TBL_EXPT_SHEET_ROW WHERE MONTH_ID = @MID AND EXPTNO IS NOT NULL);' % u)
        L.append('    IF EXISTS (SELECT 1 FROM @problems)')
        L.append('    BEGIN')
        L.append("        PRINT '*** Rollback refused - NOTHING was removed:';")
        L.append('        SELECT n AS [#], problem FROM @problems ORDER BY n;')
        L.append("        RAISERROR('Rollback refused - see the list above.', 16, 1);")
        L.append('        SET NOEXEC ON;')
        L.append('    END')
        L.append('END')
        L.append('GO')
        L.append('')
        L.append('/* ---- 2. remove: one transaction ------------------------------------------------------------ */')
        L.append('DECLARE @problems TABLE (n INT IDENTITY(1,1), problem NVARCHAR(1000));')
        L.append('DECLARE @MID INT, @N INT, @MSG NVARCHAR(2048), @L INT, @R INT, @G INT, @T INT, @M INT;')
        L.append("IF UPPER(DB_NAME()) <> 'LINKSDB20' BEGIN RAISERROR('Not on LinksDB20.', 16, 1); RETURN; END")
        L.append('BEGIN TRY')
        L.append('    BEGIN TRANSACTION;')
        L.append('    SELECT @MID = MIN(ID) FROM dbo.TBL_EXPT_SHEET_MONTH WITH (UPDLOCK, HOLDLOCK)')
        L.append('     WHERE CMPCODE = %s AND CITYCODE = %s AND SHEET_MONTH = %s AND MAKER_ID = %s;' % (sql_a(self.cmp), sql_a(self.city), sql_a(self.ym), u))
        L.append("    IF @MID IS NULL RAISERROR('The imported month is not there (any more).', 16, 1);")
        L.append('    SELECT @N = COUNT(*) FROM dbo.TBL_EXPT_SHEET_ROW WITH (UPDLOCK, HOLDLOCK) WHERE MONTH_ID = @MID;')
        L.append('    SELECT @N = COUNT(*) FROM dbo.TBL_EXPT_SHEET_GRIDROW WITH (UPDLOCK, HOLDLOCK) WHERE MONTH_ID = @MID;')
        self.pristine_checks(L, '@MID', 'Month ' + self.ym)
        L.append('    IF EXISTS (SELECT 1 FROM @problems)')
        L.append('    BEGIN')
        L.append("        SELECT @MSG = STUFF((SELECT N' | ' + problem FROM @problems ORDER BY n FOR XML PATH(''), TYPE).value('.', 'NVARCHAR(MAX)'), 1, 3, N'');")
        L.append("        RAISERROR(N'%s', 16, 1, @MSG);")
        L.append('    END')
        L.append("    DELETE L FROM dbo.TBL_EXPT_SHEET_LOG L")
        L.append("     WHERE (L.ROW_KIND = 'M' AND L.ROW_ID = @MID)")
        L.append("        OR (L.ROW_KIND = 'W' AND L.ROW_ID IN (SELECT ID FROM dbo.TBL_EXPT_SHEET_ROW WHERE MONTH_ID = @MID))")
        L.append("        OR (L.ROW_KIND = 'G' AND L.ROW_ID IN (SELECT ID FROM dbo.TBL_EXPT_SHEET_GRIDROW WHERE MONTH_ID = @MID));")
        L.append('    SET @L = @@ROWCOUNT;')
        L.append('    DELETE FROM dbo.TBL_EXPT_SHEET_ROW WHERE MONTH_ID = @MID;')
        L.append('    SET @R = @@ROWCOUNT;')
        L.append('    DELETE FROM dbo.TBL_EXPT_SHEET_GRIDROW WHERE MONTH_ID = @MID;')
        L.append('    SET @G = @@ROWCOUNT;')
        L.append('    DELETE FROM dbo.TBL_EXPT_SHEET_TABMETA WHERE MONTH_ID = @MID;')
        L.append('    SET @T = @@ROWCOUNT;')
        L.append('    DELETE FROM dbo.TBL_EXPT_SHEET_MONTH WHERE ID = @MID;')
        L.append('    SET @M = @@ROWCOUNT;')
        L.append('    COMMIT TRANSACTION;')
        L.append("    PRINT 'Removed month ID ' + CONVERT(VARCHAR(10), @MID) + ': ' + CONVERT(VARCHAR(10), @R) + ' WO rows, ' + CONVERT(VARCHAR(10), @G) + ' grid rows, '")
        L.append("          + CONVERT(VARCHAR(10), @T) + ' tab layouts, ' + CONVERT(VARCHAR(10), @L) + ' history entries, ' + CONVERT(VARCHAR(10), @M) + ' month.';")
        L.append('END TRY')
        L.append('BEGIN CATCH')
        L.append('    IF @@TRANCOUNT > 0 ROLLBACK TRANSACTION;')
        L.append('    SET @MSG = ERROR_MESSAGE();')
        L.append("    RAISERROR(N'*** Rollback failed and was undone - nothing was removed: %s', 16, 1, @MSG);")
        L.append('END CATCH')
        L.append('GO')
        L.append('SET NOEXEC OFF;')
        L.append('GO')
        L.append('')
        L.append('/* ---- 3. verify (read-only) ----------------------------------------------------------------- */')
        L.append("IF OBJECT_ID('dbo.TBL_EXPT_SHEET_MONTH', 'U') IS NOT NULL")
        L.append("    SELECT %s AS SHEET_MONTH, COUNT(*) AS MONTHS_LEFT, CASE WHEN COUNT(*) = 0 THEN 'removed' ELSE 'still there' END AS STATE"
                 % sql_a(self.ym))
        L.append('      FROM dbo.TBL_EXPT_SHEET_MONTH WHERE CMPCODE = %s AND CITYCODE = %s AND SHEET_MONTH = %s;' % (sql_a(self.cmp), sql_a(self.city), sql_a(self.ym)))
        L.append('GO')
        return L

    # -- the read-only FF JOB check -----------------------------------------------------------------------------------
    def check_script(self):
        L = []
        L.append('-- %s - READ-ONLY (three SELECT statements, nothing else).' % self.names['check'])
        L.append('-- Which FF JOB values of the workbook would link to a Links job at import time: EXPT_MASTER.EXPTNO of')
        L.append("-- CMPCODE %s / CITYCODE %s, not cancelled; then the matches whose sheet WO is not the job's invoice no." % (self.cmp, self.city))
        L.append('-- Generated with %s; holds job numbers - keep it out of git.' % self.names['import'])
        ffv = self.ff_values()
        if not ffv:
            L.append("SELECT 'no FF JOB values in the workbook' AS RESULT;")
            return L
        cte = ('WITH V (TAB_CODE, SHEET_ROW, FF_JOB, FF_KEY, WO) AS (SELECT * FROM (VALUES\n'
               + ',\n'.join('        ' + v for v in ffv)
               + '\n    ) X (TAB_CODE, SHEET_ROW, FF_JOB, FF_KEY, WO)),\n'
               + '     R AS (SELECT V.*, ' + self.match_case('V.FF_KEY', sql_a(self.cmp), sql_a(self.city)) + ' AS RESULT FROM V)')
        L.append(cte)
        L.append("SELECT RESULT, COUNT(*) AS SHEET_ROWS, COUNT(DISTINCT NULLIF(FF_KEY, '')) AS DISTINCT_JOBS FROM R GROUP BY RESULT ORDER BY RESULT;")
        L.append(cte)
        L.append("SELECT TAB_CODE, SHEET_ROW, FF_JOB, RESULT FROM R WHERE RESULT <> 'match' ORDER BY TAB_CODE, SHEET_ROW;")
        L.append('WITH V (TAB_CODE, SHEET_ROW, FF_JOB, FF_KEY, WO) AS (SELECT * FROM (VALUES')
        L.append(',\n'.join('        ' + v for v in ffv))
        L.append('    ) X (TAB_CODE, SHEET_ROW, FF_JOB, FF_KEY, WO))')
        L.append(self.wo_check_select(sql_a(self.cmp), sql_a(self.city)) + ';')
        return L

    def wo_check_select(self, cmp_expr, city_expr):
        def norm(x):
            return "UPPER(REPLACE(REPLACE(REPLACE(LTRIM(RTRIM(ISNULL(%s, N''))), N'/', N''), N'-', N''), N' ', N''))" % x
        return ("SELECT V.TAB_CODE, V.SHEET_ROW, V.FF_JOB, V.WO AS SHEET_WO, LTRIM(RTRIM(E.EXPT_INVNO)) AS JOB_INVOICE_NO,\n"
                "       CASE WHEN %s = %s THEN 'same invoice, other spelling' ELSE 'WO is not the job''s invoice no.' END AS NOTE\n"
                "  FROM V JOIN dbo.EXPT_MASTER E ON E.EXPTNO = V.FF_KEY AND LTRIM(RTRIM(ISNULL(E.CMPCODE, ''))) = %s\n"
                "       AND LTRIM(RTRIM(ISNULL(E.CITYCODE, ''))) = %s AND ISNULL(E.EXPTCANCEL, '') <> 'Y'\n"
                " WHERE LTRIM(RTRIM(ISNULL(E.EXPT_INVNO, ''))) <> LTRIM(RTRIM(V.WO))\n"
                " ORDER BY V.TAB_CODE, V.SHEET_ROW") % (norm('E.EXPT_INVNO'), norm('V.WO'), cmp_expr, city_expr)


# ======================================================================================================================
# main
# ======================================================================================================================
def inside_git(path):
    p = os.path.abspath(path)
    while True:
        if os.path.exists(os.path.join(p, '.git')):
            return p
        q = os.path.dirname(p)
        if q == p:
            return None
        p = q


def write_lines(path, lines):
    text = '\r\n'.join(lines) + '\r\n'
    bad = [i for i, ch in enumerate(text) if ord(ch) > 126 or (ord(ch) < 32 and ch not in '\r\n')]
    if bad:
        raise AssertionError('%s would not be plain ASCII (offset %d)' % (path, bad[0]))
    with open(path, 'w', encoding='ascii', newline='') as f:
        f.write(text)


def main(argv=None):
    ap = argparse.ArgumentParser(description='Bajaj monthly workbook -> Export > Sheet view import script (LinksDB20).')
    ap.add_argument('workbook', nargs='?', help='the .xlsx downloaded from Google Sheets')
    ap.add_argument('--month', help='sheet month YYYY-MM')
    ap.add_argument('--title', help='month title (TBL_EXPT_SHEET_MONTH.TITLE, max 100)')
    ap.add_argument('--out-dir', help='where the scripts go - outside any git work tree')
    ap.add_argument('--cmpcode', default='30', help='company the Sheet view logs in with (default 30 = LINKS CARGO AGENCIES PVT. LTD.)')
    ap.add_argument('--citycode', default='MUM', help='branch the Sheet view logs in with (default MUM)')
    ap.add_argument('--registry', help='sheet-registry.ts (default: ../../src/app/export/sheet/sheet-registry.ts from this file)')
    ap.add_argument('--skip-tab', action='append', help='a workbook tab to leave out (repeatable)')
    ap.add_argument('--date', help='date in the file names, YYYYMMDD (default today)')
    ap.add_argument('--allow-repo-output', action='store_true', help='allow --out-dir inside a git work tree (never for real data)')
    ap.add_argument('--selftest', action='store_true', help='check the value rules and exit')
    a = ap.parse_args(argv)
    if a.selftest:
        return selftest()
    errs = []
    if not a.workbook or not os.path.isfile(a.workbook):
        errs.append('workbook not found')
    if not a.month or not re.fullmatch(r'[12][0-9]{3}-(0[1-9]|1[0-2])', a.month):
        errs.append('--month must be YYYY-MM')
    if not a.title or not a.title.strip() or len(a.title.strip()) > 100:
        errs.append('--title is required, at most 100 characters')
    if not a.out_dir:
        errs.append('--out-dir is required')
    if not re.fullmatch(r'[A-Za-z0-9]{1,6}', a.cmpcode or ''):
        errs.append('--cmpcode: 1-6 letters / digits')
    if not re.fullmatch(r'[A-Za-z0-9]{1,3}', a.citycode or ''):
        errs.append('--citycode: 1-3 letters / digits')
    if a.date and not re.fullmatch(r'20[0-9]{6}', a.date):
        errs.append('--date must be YYYYMMDD')
    if errs:
        ap.error('; '.join(errs))
    a.title = a.title.strip()
    os.makedirs(a.out_dir, exist_ok=True)
    repo = inside_git(a.out_dir)
    if repo and not a.allow_repo_output:
        ap.error('--out-dir is inside the git work tree %s - the scripts hold customer data; use a folder outside it' % repo)
    reg_path = a.registry or os.path.normpath(os.path.join(os.path.dirname(os.path.abspath(__file__)), '..', '..', 'src', 'app', 'export', 'sheet', 'sheet-registry.ts'))
    reg = Registry(reg_path)
    rep = Report()
    wb_hash = hashlib.sha256(open(a.workbook, 'rb').read()).hexdigest()
    imp = Import(a, reg, rep)
    rep.add('Sheet view import - %s' % VERSION)
    rep.add('Workbook : %s (SHA-256 %s)' % (os.path.basename(a.workbook), wb_hash))
    rep.add('Target   : LinksDB20, CMPCODE %s / CITYCODE %s, month %s, title "%s"' % (a.cmpcode, a.citycode, a.month, a.title))
    rep.add('Registry : %s' % reg_path)
    rep.add('')
    imp.run()
    rep.add('')
    rep.add('Statuses (row fill -> Colour CODE): ' + (', '.join('%s %d' % kv for kv in sorted(imp.status_count.items())) or 'none'))
    rep.add('Rows without a status: %d' % sum(1 for x in imp.wo if not x['status']))
    if imp.unknown_colours:
        rep.add('Unknown row colours (no Colour CODE status; STATUS_CODE left empty, the fill kept per cell in CELL_FMT):')
        for c, rows in sorted(imp.unknown_colours.items()):
            rep.add('  %s: %d row(s): %s' % (c, len(rows), ', '.join(rows)))
    else:
        rep.add('Unknown row colours: none')
    ffs = imp.ff
    keys = [k for _, _, _, k, _ in ffs if k]
    rep.add('FF JOB: %d rows carry a value, %d of them job-like (%d distinct); matched to EXPT_MASTER at run time.'
            % (len(ffs), len(keys), len(set(keys))))
    for tab, r, text, key, _ in ffs:
        if not key:
            rep.add('  not a job number: %s row %d: "%s" (stays FF JOB text)' % (tab, r, ascii_text(text)))
    rep.add('')
    for w in rep.warnings:
        rep.add('NOTE: ' + w)
    stamp = a.date or datetime.date.today().strftime('%Y%m%d')
    ym = a.month.replace('-', '')
    names = {'import': 'IMPORT_EXPORT_SHEET_%s_LINKSDB20_%s.sql' % (ym, stamp),
             'rollback': 'ROLLBACK_IMPORT_EXPORT_SHEET_%s_LINKSDB20_%s.sql' % (ym, stamp),
             'check': 'CHECK_FFJOB_EXPORT_SHEET_%s.sql' % ym,
             'report': 'IMPORT_EXPORT_SHEET_%s_REPORT.txt' % ym}
    if rep.refusals:
        rep.add('')
        rep.add('REFUSED - no script written. Fix the workbook (or the registry) and run again:')
        for r in rep.refusals:
            rep.add('  * ' + r)
        out = '\n'.join(rep.lines)
        print(out)
        write_lines(os.path.join(a.out_dir, names['report']), [ascii_text(l) for l in rep.lines])
        return 2
    sql = Sql(imp, a, names, wb_hash)
    write_lines(os.path.join(a.out_dir, names['import']), sql.import_script())
    write_lines(os.path.join(a.out_dir, names['rollback']), sql.rollback_script())
    write_lines(os.path.join(a.out_dir, names['check']), sql.check_script())
    rep.add('Written to %s:' % os.path.abspath(a.out_dir))
    for k in ('import', 'rollback', 'check', 'report'):
        rep.add('  ' + names[k])
    write_lines(os.path.join(a.out_dir, names['report']), [ascii_text(l) for l in rep.lines])
    print('\n'.join(rep.lines))
    return 0


# ======================================================================================================================
# self-test of the value rules
# ======================================================================================================================
def selftest():
    import json
    ok = True

    def check(name, got, want):
        nonlocal ok
        if got != want:
            ok = False
            print('FAIL %s: got %r, want %r' % (name, got, want))

    # SQL collation order of the 52 keys, as LinksDB20 returned it on 2026-10-21 (SELECT ... ORDER BY K)
    sql_order = ('AGENT BL_DT BL_HANDOVER_DT BL_NO BOOKING_NO BUFFER_YARD CARGO_TYPE CLEARANCE_POINT CNTR_DISPATCH_DT CNTR_GATEDIN_DT '
                 'CNTR_REPORT_NS_DT CNTR_TYPE CONSIGNEE CONT CONTAINER_NO COUNTRY COURIER_DT CURRENT_ETD DO_ETD DO_ETD2 DO_GIVEN_DT '
                 'DO_RECEIVE_DT EDOC_STATUS ETA_DEST FF_JOB FINAL_VSL_SOB FOR_HBL GATE_CUTOFF_DT GATE_OPEN_DT HAZ LC_DATE LC_NO '
                 'OPEN_ORDER PICKUP_DT PLANT PO_NO POL_GATE PORT QTY REMARK1 SB_DT SB_NO SI_CUTOFF_DT SI_SUBMITTED_DT SLINE '
                 'SLINE_PAYMENT_STATUS STUFFING_DT TRANSPORTER VEH VGM_SUBMITTED_DT VSL_NAME WO').split()
    check('key order', sorted(sql_order[::-1], key=sql_key_order), sql_order)
    check('letters', sorted(['B', 'AA', 'A', 'AB'], key=sql_key_order), ['A', 'AA', 'AB', 'B'])
    # dates (USP_EXPT_SHEET_SAVE_APPLY)
    for v, w in (('14-10-2026', '14-10-2026'), ('14/10/2026', '14-10-2026'), (' 14.10.2026 ', '14-10-2026'), ('2026-10-14', '14-10-2026'),
                 ('14-10/2026', '14-10/2026'), ('26-08/27-08', '26-08/27-08'), ('31-02-2026', '31-02-2026'), ('Will be updated', 'Will be updated'),
                 ('1-10-2026', '1-10-2026'), ('29-02-2028', '29-02-2028'), ('01-01-1700', '01-01-1700')):
        check('date %r' % v, normalize_date_text(v), w)
    # numbers
    for v, w in ((6.0, '6'), (301102610004227.0, '301102610004227'), (1.5, '1.5'), (0.1, '0.1'), (1e-7, '0.0000001'), (12, '12'), (True, 'TRUE')):
        check('num %r' % v, num_text(v), w)
    # Json.NET Formatting.None
    check('json', nj({'A': 'a"b\\c\nd\r\te\x01 é\U0001F600', 'B': 'x'}),
          '{"A":"a\\"b\\\\c\\nd\\r\\te\\u0001\\u2028é\U0001F600","B":"x"}')
    check('json types', nj([1, True, None, {'k': []}]), '[1,true,null,{"k":[]}]')
    check('guard', guard_json('{"A":"x|;amp;|y"}'), '{"A":"x\\u007C;amp;\\u007Cy"}')
    check('json roundtrip', json.loads(nj({'A': 'q"\n '})), {'A': 'q"\n '})
    # XmlSafe
    check('xmlsafe', xml_safe('a\r\nb\rc\x00d￾'), ('a\nb\ncd', 2, 2))
    # CELL_FMT canonical
    f = cell_fmt_json({'WO': {'bg': '#FF00FF', 'b': 1}, 'VEH': {'fg': '#305496'}, 'PORT': {}}, sql_key_order)
    check('cell_fmt', f, '{"VEH":{"fg":"#305496"},"WO":{"bg":"#ff00ff","b":1}}')
    check('cell_fmt canonical', cell_fmt_is_canonical(f), True)
    check('cell_fmt empty', cell_fmt_json({'WO': {}}, sql_key_order), None)
    check('not canonical', cell_fmt_is_canonical('{"WO": {"bg":"#ff00ff"}}'), False)
    check('not canonical 2', cell_fmt_is_canonical('{"WO":{"color":"red"}}'), False)
    # SQL literals
    check('sql plain', sql_n("O'Neil"), "N'O''Neil'")
    check('sql nl', sql_n('a\nb'), "CONVERT(NVARCHAR(MAX), N'a') + NCHAR(10) + N'b'")
    check('sql astral', sql_n('\U0001F600'), 'CONVERT(NVARCHAR(MAX), NCHAR(55357)) + NCHAR(56832)')
    check('sql null', sql_n(None), 'NULL')
    long = 'x' * 2500
    check('sql chunk', sql_n(long).count(" + N'"), 2)
    # TS literal parsing
    check('ts', json.loads(_ts_to_json("{ a: 'it\\'s \"q\"', b: [1, 2,], c: \"40'HC\" , /* x */ d: true, }")),
          {'a': 'it\'s "q"', 'b': [1, 2], 'c': "40'HC", 'd': True})
    check('ts2', json.loads(_ts_to_json("[{ code: 'NIGERIA', title: 'Nigeria' }]")), [{'code': 'NIGERIA', 'title': 'Nigeria'}])
    # colours
    class C:
        pass
    c = C(); c.type = 'theme'; c.theme = 9; c.tint = 0.0
    col = Colours.__new__(Colours)
    col.theme = ['FFFFFF', '000000', 'FFFFFF', '000000', '4285F4', 'EA4335', 'FBBC04', '34A853', 'FF6D01', '46BDC6', '1155CC', '1155CC']
    col.indexed = []
    check('theme9', col.resolve(c), '#46bdc6')
    c.tint = 0.7999816888943144
    check('tint', col.resolve(c), '#daf2f4')
    print('selftest ' + ('OK' if ok else 'FAILED'))
    return 0 if ok else 1


if __name__ == '__main__':
    sys.exit(main())
